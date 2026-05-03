#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import time
from copy import copy
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Color, Font, PatternFill

from value_investment_market_value import (
    _detect_report_currency,
    _infer_common_years,
    _normalize_ccy,
    _normalize_code,
    _quote_currency_from_code,
    _report_currency_marked,
    _resolve_fx_rate,
    _sheet_saved_report_currency,
    build_sheet_rows,
    parse_number,
    write_sheet,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FINANCE_ROOT = PROJECT_ROOT / "data" / "finance_report_web"
ROW_FONT_RED = "FF0000"
ROW_FONT_GRAY = "808080"
ROW_FONT_GREEN = "008000"
NO_FILL = PatternFill(fill_type=None)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run valuation for codes in value_analysis_by_code workbook")
    parser.add_argument(
        "--workbook",
        default=str(PROJECT_ROOT / "data" / "value_analysis_by_code.xlsx"),
        help="Target workbook path (read codes and write D/E + calc sheets)",
    )
    parser.add_argument(
        "--home-sheet",
        default="",
        help="Sheet containing company/code list (default: first sheet)",
    )
    parser.add_argument(
        "--fx-workbook",
        default=str(PROJECT_ROOT / "data" / "Value_Investment.xlsx"),
        help="Workbook used to resolve FX matrix from 统计 sheet when needed",
    )
    parser.add_argument(
        "--years",
        type=int,
        default=5,
        help="Latest N common fiscal years to use",
    )
    parser.add_argument(
        "--yahoo-timeout",
        type=int,
        default=10,
        help="Yahoo timeout seconds for shares fallback",
    )
    parser.add_argument(
        "--no-download-missing",
        action="store_true",
        help="Disable automatic download for missing local financial data",
    )
    return parser.parse_args()


def calc_sheet_name(code: str) -> str:
    base = code.replace(".", "_")
    base = re.sub(r"[:\\\\/?*\\[\\]]", "_", base)
    return base[:31]


def _financial_paths(code: str) -> List[Path]:
    return [
        FINANCE_ROOT / f"{code}_income_year.csv",
        FINANCE_ROOT / f"{code}_balance_year.csv",
        FINANCE_ROOT / f"{code}_cash_year.csv",
    ]


def _info_path(code: str) -> Path:
    return FINANCE_ROOT / f"{code}_info.csv"


def has_local_financial_data(code: str) -> bool:
    return all(p.exists() for p in _financial_paths(code))


def _downloader_python_bin() -> str:
    py = Path("/Users/jacky/miniconda3/bin/python3")
    if py.exists():
        return str(py)
    return sys.executable


def download_financial_data(code: str, retries: int = 2, timeout_sec: int = 240) -> bool:
    script = PROJECT_ROOT / "scripts" / "download_financial_data.py"
    if not script.exists():
        print(f"[{code}] downloader script not found: {script}")
        return False

    py_bin = _downloader_python_bin()
    env = os.environ.copy()
    env.setdefault("MPLCONFIGDIR", "/tmp/mplconfig")
    env.setdefault("PYTHONUNBUFFERED", "1")

    for attempt in range(1, retries + 1):
        try:
            cmd = [py_bin, str(script), code]
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout_sec, env=env)
            if proc.returncode != 0:
                tail = (proc.stderr or proc.stdout or "").strip().splitlines()[-3:]
                print(f"[{code}] download failed (attempt {attempt}/{retries}), rc={proc.returncode}: {' | '.join(tail)}")
                time.sleep(min(1.5 * attempt, 3))
                continue
            if has_local_financial_data(code):
                return True
        except subprocess.TimeoutExpired:
            print(f"[{code}] download timeout (attempt {attempt}/{retries}, {timeout_sec}s)")
            time.sleep(min(1.5 * attempt, 3))
        except Exception as exc:
            print(f"[{code}] download exception (attempt {attempt}/{retries}): {exc}")
            time.sleep(min(1.5 * attempt, 3))
    return has_local_financial_data(code)


def load_financial_frames(code: str) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    income_df = pd.read_csv(FINANCE_ROOT / f"{code}_income_year.csv", index_col=0, encoding="utf-8-sig")
    balance_df = pd.read_csv(FINANCE_ROOT / f"{code}_balance_year.csv", index_col=0, encoding="utf-8-sig")
    cash_df = pd.read_csv(FINANCE_ROOT / f"{code}_cash_year.csv", index_col=0, encoding="utf-8-sig")
    if _info_path(code).exists():
        info_df = pd.read_csv(_info_path(code), index_col=0, encoding="utf-8-sig")
    else:
        # info.csv is optional: shares can still be resolved from Yahoo fallback.
        info_df = pd.DataFrame()
    return income_df, balance_df, cash_df, info_df


def _extract_periods(df: pd.DataFrame) -> set:
    periods = set()
    for c in df.columns:
        m = re.match(r"^(\d{4}/\d{2}/\d{2})", str(c))
        if m:
            periods.add(m.group(1))
    return periods


def latest_common_report_period(income_df: pd.DataFrame, balance_df: pd.DataFrame, cash_df: pd.DataFrame) -> Tuple[str, int]:
    common = _extract_periods(income_df) & _extract_periods(balance_df) & _extract_periods(cash_df)
    if not common:
        return "", 0
    latest = sorted(common)[-1]
    try:
        return latest, int(latest[:4])
    except Exception:
        return latest, 0


def fallback_report_currency(code: str) -> str:
    prefix = code.split(".", 1)[0]
    return {
        "US": "USD",
        "JP": "JPY",
        "SH": "CNY",
        "SZ": "CNY",
        "HK": "CNY",
    }.get(prefix, "CNY")


def _yahoo_symbol_from_code(code: str) -> Optional[str]:
    c = _normalize_code(code)
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "HK":
        if ticker.isdigit():
            return f"{int(ticker):04d}.HK"
        return f"{ticker}.HK"
    if market == "SH":
        return f"{ticker}.SS"
    if market == "SZ":
        return f"{ticker}.SZ"
    if market == "US":
        return ticker
    if market == "JP":
        return f"{ticker}.T"
    return None


def _parse_price(value: object) -> Optional[float]:
    s = str(value or "").strip()
    if not s or s == "--":
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", s.replace(",", ""))
    if not m:
        return None
    try:
        return float(m.group(0))
    except ValueError:
        return None


def _code_to_tencent_symbol(code: str) -> Optional[str]:
    c = _normalize_code(code)
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "HK":
        return f"hk{ticker}"
    if market == "SH":
        return f"sh{ticker}"
    if market == "SZ":
        return f"sz{ticker}"
    if market == "US":
        return f"us{ticker}"
    if market == "JP":
        return f"jp{ticker}"
    return None


def _fetch_prices_tencent(codes: List[str], timeout: int = 8) -> Dict[str, float]:
    symbol_to_code: Dict[str, str] = {}
    symbols: List[str] = []
    seen: set = set()
    for code in codes:
        symbol = _code_to_tencent_symbol(code)
        if symbol:
            key = symbol.lower()
            symbol_to_code[key] = _normalize_code(code)
            if key not in seen:
                symbols.append(symbol)
                seen.add(key)
    if not symbols:
        return {}

    out: Dict[str, float] = {}
    headers = {"User-Agent": "Mozilla/5.0"}
    for i in range(0, len(symbols), 60):
        chunk = symbols[i : i + 60]
        url = "https://qt.gtimg.cn/q=" + ",".join(chunk)
        try:
            resp = requests.get(url, headers=headers, timeout=timeout)
            resp.raise_for_status()
            text = resp.text
        except Exception:
            continue

        for m in re.finditer(r'v_([^=]+)="([^"]*)";', text):
            sym = m.group(1).strip().lower()
            payload = m.group(2)
            code = symbol_to_code.get(sym)
            if not code:
                continue
            parts = payload.split("~")
            if len(parts) < 4:
                continue
            # Prefer previous close (index 4). Fall back to live trade price (index 3) only if needed.
            price = _parse_price(parts[4]) if len(parts) > 4 else None
            if price is None:
                price = _parse_price(parts[3])
            if price is None:
                numeric = [_parse_price(x) for x in parts[1:12]]
                numeric = [x for x in numeric if x is not None and x > 0]
                price = numeric[0] if numeric else None
            if price is not None and price > 0:
                out[code] = float(price)
    return out


def _fetch_dividend_yield_tencent(codes: List[str], timeout: int = 8) -> Dict[str, float]:
    symbol_to_code: Dict[str, str] = {}
    symbols: List[str] = []
    seen: set = set()
    for code in codes:
        symbol = _code_to_tencent_symbol(code)
        if symbol:
            key = symbol.lower()
            symbol_to_code[key] = _normalize_code(code)
            if key not in seen:
                symbols.append(symbol)
                seen.add(key)
    if not symbols:
        return {}

    out: Dict[str, float] = {}
    headers = {"User-Agent": "Mozilla/5.0"}
    for i in range(0, len(symbols), 60):
        chunk = symbols[i : i + 60]
        url = "https://qt.gtimg.cn/q=" + ",".join(chunk)
        try:
            resp = requests.get(url, headers=headers, timeout=timeout)
            resp.raise_for_status()
            text = resp.text
        except Exception:
            continue

        for m in re.finditer(r'v_([^=]+)="([^"]*)";', text):
            sym = m.group(1).strip().lower()
            payload = m.group(2)
            code = symbol_to_code.get(sym)
            if not code:
                continue
            parts = payload.split("~")
            if not parts:
                continue
            market = code.split(".", 1)[0]
            idx = None
            # Empirical mapping validated by current payloads:
            # HK: idx 47, US: idx 52, SH/SZ: idx 43 (all in percent units).
            if market == "HK":
                idx = 47
            elif market == "US":
                idx = 52
            elif market in {"SH", "SZ"}:
                idx = 43
            if idx is None or idx >= len(parts):
                continue
            pct = _parse_price(parts[idx])
            if pct is None:
                continue
            if pct < 0 or pct > 100:
                continue
            out[code] = float(pct) / 100.0
    return out


def _fetch_prices_yahoo(codes: List[str], timeout: int = 12) -> Dict[str, float]:
    sym_to_code: Dict[str, str] = {}
    for code in codes:
        sym = _yahoo_symbol_from_code(code)
        if sym:
            sym_to_code[sym.upper()] = code
    symbols = list(sym_to_code.keys())
    if not symbols:
        return {}

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
    }
    out: Dict[str, float] = {}

    for i in range(0, len(symbols), 50):
        chunk = symbols[i : i + 50]
        url = "https://query1.finance.yahoo.com/v7/finance/quote"
        params = {"symbols": ",".join(chunk)}
        try:
            resp = requests.get(url, params=params, headers=headers, timeout=timeout)
            if resp.status_code != 200:
                continue
            data = resp.json()
            result = data.get("quoteResponse", {}).get("result", [])
            for item in result:
                sym = str(item.get("symbol", "")).upper().strip()
                code = sym_to_code.get(sym)
                if not code:
                    continue
                # Prefer previous close so the daily push shows "recent close" rather than intraday last trade.
                for key in ("regularMarketPreviousClose", "previousClose", "chartPreviousClose", "regularMarketPrice"):
                    price = item.get(key)
                    if isinstance(price, (int, float)) and price > 0:
                        out[code] = float(price)
                        break
        except Exception:
            continue
    return out


def _fetch_prices_yahoo_jp_page(codes: List[str], timeout: int = 12) -> Dict[str, float]:
    out: Dict[str, float] = {}
    headers = {"User-Agent": "Mozilla/5.0"}
    for code in codes:
        c = _normalize_code(code)
        if not c.startswith("JP.") or "." not in c:
            continue
        ticker = c.split(".", 1)[1].strip()
        if not ticker:
            continue
        url = f"https://finance.yahoo.co.jp/quote/{ticker}.T"
        try:
            resp = requests.get(url, headers=headers, timeout=timeout)
            resp.raise_for_status()
            html = resp.text or ""
        except Exception:
            continue

        prev_close = None
        m_prev = re.search(
            r"前日終値.*?DataListItem__value__[^>]*>([0-9,]+(?:\.[0-9]+)?)</",
            html,
            re.S,
        )
        if m_prev:
            prev_close = _parse_price(m_prev.group(1))

        pattern_json = re.compile(
            rf'"codeWithMarketExtension":"{re.escape(ticker)}\.T".{{0,1500}}?"price":"([0-9,]+(?:\.[0-9]+)?)"',
            re.S,
        )
        m = pattern_json.search(html)
        price = _parse_price(m.group(1)) if m else None
        if price is None:
            m2 = re.search(
                r'PriceBoard__price__[^>]*>.*?StyledNumber__value__[^>]*>([0-9,]+(?:\.[0-9]+)?)<',
                html,
                re.S,
            )
            price = _parse_price(m2.group(1)) if m2 else None
        chosen = prev_close if prev_close is not None and prev_close > 0 else price
        if chosen is not None and chosen > 0:
            out[c] = float(chosen)
    return out


def _fetch_dividend_yield_yahoo_jp_page(codes: List[str], timeout: int = 12) -> Dict[str, float]:
    out: Dict[str, float] = {}
    headers = {"User-Agent": "Mozilla/5.0"}
    for code in codes:
        c = _normalize_code(code)
        if not c.startswith("JP.") or "." not in c:
            continue
        ticker = c.split(".", 1)[1].strip()
        if not ticker:
            continue
        url = f"https://finance.yahoo.co.jp/quote/{ticker}.T"
        try:
            resp = requests.get(url, headers=headers, timeout=timeout)
            resp.raise_for_status()
            html = resp.text or ""
        except Exception:
            continue
        # Example block:
        # 配当利回り ... DataListItem__value__11kV">1.57</span><span ...>%</span>
        m = re.search(
            r"配当利回り.*?DataListItem__value__[^>]*>([0-9]+(?:\.[0-9]+)?)</span>\s*<span[^>]*>%</span>",
            html,
            re.S,
        )
        if not m:
            continue
        pct = _parse_price(m.group(1))
        if pct is None or pct < 0 or pct > 100:
            continue
        out[c] = float(pct) / 100.0
    return out


def _code_to_stooq_symbol(code: str) -> Optional[str]:
    c = _normalize_code(code)
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "US":
        return f"{ticker.lower()}.us"
    if market == "JP":
        return f"{ticker.lower()}.jp"
    return None


def _fetch_prices_stooq(codes: List[str], timeout: int = 10) -> Dict[str, float]:
    symbol_to_code: Dict[str, str] = {}
    for code in codes:
        symbol = _code_to_stooq_symbol(code)
        if symbol:
            symbol_to_code[symbol] = code
    symbols = list(symbol_to_code.keys())
    if not symbols:
        return {}

    out: Dict[str, float] = {}
    headers = {"User-Agent": "Mozilla/5.0"}
    for symbol in symbols:
        try:
            resp = requests.get(
                "https://stooq.com/q/l/",
                params={"s": symbol, "i": "d"},
                headers=headers,
                timeout=timeout,
            )
            resp.raise_for_status()
            text = (resp.text or "").strip()
            if not text:
                continue
            parts = [x.strip() for x in text.splitlines()[0].split(",")]
            if len(parts) < 7:
                continue
            price = _parse_price(parts[6])
            if price is not None and price > 0:
                out[symbol_to_code[symbol]] = float(price)
        except Exception:
            continue
    return out


def _build_quote_cache(codes: List[str], timeout: int = 12) -> Tuple[Dict[str, float], Dict[str, str]]:
    quote_prices: Dict[str, float] = {}
    quote_sources: Dict[str, str] = {}
    uniq_codes = sorted(set(_normalize_code(x) for x in codes if str(x).strip()))

    yahoo_prices = _fetch_prices_yahoo(uniq_codes, timeout=timeout) if uniq_codes else {}
    for c, p in yahoo_prices.items():
        quote_prices[c] = p
        quote_sources[c] = "yahoo"
    print(f"Yahoo quotes fetched: {len(yahoo_prices)}/{len(uniq_codes)}")

    remain = [c for c in uniq_codes if c not in quote_prices]
    remain_jp = [c for c in remain if c.startswith("JP.")]
    yahoo_jp_prices = _fetch_prices_yahoo_jp_page(remain_jp, timeout=min(timeout, 12)) if remain_jp else {}
    for c, p in yahoo_jp_prices.items():
        quote_prices[c] = p
        quote_sources[c] = "yahoo_jp_page"
    print(f"Yahoo JP page quotes fetched: {len(yahoo_jp_prices)}/{len(remain_jp)}")

    remain = [c for c in uniq_codes if c not in quote_prices]
    remain_stooq = [c for c in remain if c.startswith(("US.", "JP."))]
    stooq_prices = _fetch_prices_stooq(remain_stooq, timeout=min(timeout, 10)) if remain_stooq else {}
    for c, p in stooq_prices.items():
        quote_prices[c] = p
        quote_sources[c] = "stooq"
    print(f"Stooq quotes fetched: {len(stooq_prices)}/{len(remain_stooq)}")

    remain = [c for c in uniq_codes if c not in quote_prices]
    tencent_prices = _fetch_prices_tencent(remain, timeout=min(timeout, 10)) if remain else {}
    for c, p in tencent_prices.items():
        quote_prices[c] = p
        quote_sources[c] = "tencent"
    print(f"Tencent quotes fetched: {len(tencent_prices)}/{len(remain)}")
    return quote_prices, quote_sources


def _dividend_yield_from_info(info_df: pd.DataFrame) -> Optional[float]:
    if info_df is None or info_df.empty:
        return None
    keys = ["股息率", "股息率TTM", "Dividend Yield"]
    for key in keys:
        if key not in info_df.index:
            continue
        row = info_df.loc[key]
        vals = row.tolist() if isinstance(row, pd.Series) else [row]
        for v in vals:
            s = str(v or "").strip()
            if not s or s == "--":
                continue
            if "%" in s:
                n = _parse_price(s)
                if n is not None:
                    return float(n) / 100.0
            n = _parse_price(s)
            if n is None:
                continue
            # Heuristic: >1 likely percentage units; <=1 likely already ratio.
            if n > 1:
                return float(n) / 100.0
            return float(n)
    return None


def _build_dividend_yield_cache(codes: List[str], timeout: int = 12) -> Dict[str, float]:
    uniq_codes = sorted(set(_normalize_code(x) for x in codes if str(x).strip()))
    out: Dict[str, float] = {}

    dy_tencent = _fetch_dividend_yield_tencent(uniq_codes, timeout=min(timeout, 10))
    out.update(dy_tencent)
    print(f"Tencent dividend yields fetched: {len(dy_tencent)}/{len(uniq_codes)}")

    remain = [c for c in uniq_codes if c not in out and c.startswith("JP.")]
    dy_jp = _fetch_dividend_yield_yahoo_jp_page(remain, timeout=min(timeout, 12)) if remain else {}
    out.update(dy_jp)
    print(f"Yahoo JP dividend yields fetched: {len(dy_jp)}/{len(remain)}")
    return out


def _latest_close_from_info(info_df: pd.DataFrame) -> Optional[float]:
    if info_df is None or info_df.empty:
        return None
    # Keep this strictly aligned to "recent close" semantics.
    keys = ["收盘价", "收盤價", "昨收", "昨收价", "昨收價", "前收盘价", "前收盤價", "前收"]
    for key in keys:
        if key in info_df.index:
            row = info_df.loc[key]
            vals = row.tolist() if isinstance(row, pd.Series) else [row]
            for v in vals:
                n = parse_number(v)
                if n > 0:
                    return float(n)
    return None


def _latest_close_from_refreshed_futu_info(info_df: pd.DataFrame) -> Optional[float]:
    close = _latest_close_from_info(info_df)
    if close and close > 0:
        return close
    # Historical getStockInfoWeb() stores 昨收 into the "平均价" row label.
    if info_df is None or info_df.empty:
        return None
    if "平均价" in info_df.index:
        row = info_df.loc["平均价"]
        vals = row.tolist() if isinstance(row, pd.Series) else [row]
        for v in vals:
            n = parse_number(v)
            if n > 0:
                return float(n)
    return None


def refresh_info_via_futu(code: str) -> pd.DataFrame:
    try:
        from download_financial_data import getStockInfoWeb
    except Exception:
        return pd.DataFrame()
    try:
        getStockInfoWeb(code)
    except Exception:
        return pd.DataFrame()
    if _info_path(code).exists():
        try:
            return pd.read_csv(_info_path(code), index_col=0, encoding="utf-8-sig")
        except Exception:
            return pd.DataFrame()
    return pd.DataFrame()


def resolve_latest_close(
    code: str,
    info_df: pd.DataFrame,
    quote_cache: Optional[Dict[str, float]] = None,
    quote_source_cache: Optional[Dict[str, str]] = None,
    refresh_futu_on_missing: bool = True,
) -> Tuple[Optional[float], str]:
    c = _normalize_code(code)
    if quote_cache is not None:
        q = quote_cache.get(c)
        src = quote_source_cache.get(c, "cache") if quote_source_cache else "cache"
        if q and q > 0 and src != "tencent":
            return float(q), src
    if refresh_futu_on_missing:
        refreshed = refresh_info_via_futu(c)
        f2 = _latest_close_from_refreshed_futu_info(refreshed)
        if f2 and f2 > 0:
            return float(f2), "futu_refresh"
    f = _latest_close_from_info(info_df)
    if f and f > 0:
        return float(f), "info_cache"
    if quote_cache is not None:
        q = quote_cache.get(c)
        if q and q > 0:
            src = quote_source_cache.get(c, "cache") if quote_source_cache else "cache"
            return float(q), src
    return None, "missing"


def calc_revenue_cagr_3y(rows: List[object]) -> Optional[float]:
    rev_series = next((x for x in rows if getattr(x, "label", "") == "Revenue"), None)
    if rev_series is None:
        return None
    vals = list(getattr(rev_series, "values", []))
    if len(vals) >= 4:
        start = float(vals[-4])
        end = float(vals[-1])
        years = 3
    elif len(vals) >= 2:
        start = float(vals[0])
        end = float(vals[-1])
        years = len(vals) - 1
    else:
        return None
    if start <= 0 or end <= 0 or years <= 0:
        return None
    try:
        return (end / start) ** (1.0 / years) - 1.0
    except Exception:
        return None


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook).resolve()
    fx_workbook = Path(args.fx_workbook).resolve()

    wb = load_workbook(workbook_path, data_only=False)
    home_sheet_name = args.home_sheet.strip() or wb.sheetnames[0]
    if home_sheet_name not in wb.sheetnames:
        raise ValueError(f"home sheet not found: {home_sheet_name}")
    ws = wb[home_sheet_name]

    tasks: List[Tuple[int, str, str]] = []
    for row in range(2, ws.max_row + 1):
        company = str(ws.cell(row, 2).value or "").strip()
        raw_code = str(ws.cell(row, 3).value or "").strip()
        if not raw_code:
            continue
        code = _normalize_code(raw_code)
        tasks.append((row, company, code))
    row_to_code = {row: code for row, _, code in tasks}

    wb.close()

    wb_check = load_workbook(workbook_path, data_only=True)
    has_stats_in_target = "统计" in wb_check.sheetnames
    wb_check.close()

    results: Dict[int, Dict[str, object]] = {}
    success = 0
    failed = 0
    quote_cache, quote_source_cache = _build_quote_cache([code for _, _, code in tasks], timeout=args.yahoo_timeout)
    dividend_yield_cache = _build_dividend_yield_cache([code for _, _, code in tasks], timeout=args.yahoo_timeout)

    for row, company, code in tasks:
        print(f"\n[{row}] {company or code} {code}")
        if (not has_local_financial_data(code)) and (not args.no_download_missing):
            print(f"[{code}] local data missing, downloading...")
            ok = download_financial_data(code)
            if not ok:
                print(f"[{code}] download failed, skip.")
                results[row] = {
                    "close": None,
                    "lower": None,
                    "upper": None,
                    "price_source": "missing_data",
                    "revenue_cagr_3y": None,
                    "dividend_yield": None,
                    "price_currency": "",
                    "report_period": "N/A",
                }
                failed += 1
                continue
        elif not has_local_financial_data(code):
            print(f"[{code}] local data missing and download disabled, skip.")
            results[row] = {
                "close": None,
                "lower": None,
                "upper": None,
                "price_source": "missing_data",
                "revenue_cagr_3y": None,
                "dividend_yield": None,
                "price_currency": "",
                "report_period": "N/A",
            }
            failed += 1
            continue

        try:
            income_df, balance_df, cash_df, info_df = load_financial_frames(code)
            latest_period, latest_year = latest_common_report_period(income_df, balance_df, cash_df)
            current_year = time.localtime().tm_year
            # If report data is too old, try refresh once (common latest year <= current_year-3).
            if (latest_year > 0) and (latest_year <= current_year - 3) and (not args.no_download_missing):
                print(
                    f"[{code}] report looks stale (latest={latest_period or latest_year}), "
                    "retry download latest financials..."
                )
                if download_financial_data(code, retries=1):
                    income_df, balance_df, cash_df, info_df = load_financial_frames(code)
                    latest_period, latest_year = latest_common_report_period(income_df, balance_df, cash_df)
                    print(f"[{code}] after refresh latest report period: {latest_period or 'N/A'}")
            years = _infer_common_years(income_df, balance_df, cash_df, count=args.years)

            report_currency = _detect_report_currency(income_df, balance_df, cash_df)
            if not _report_currency_marked(income_df, balance_df, cash_df):
                prev_ccy = _sheet_saved_report_currency(workbook_path, calc_sheet_name(code))
                if not prev_ccy and fx_workbook.exists():
                    prev_ccy = _sheet_saved_report_currency(fx_workbook, f"{company}_新估值")
                report_currency = prev_ccy or fallback_report_currency(code)
            report_currency = _normalize_ccy(report_currency)

            quote_currency = _quote_currency_from_code(code)
            fx_source_book = workbook_path if has_stats_in_target else fx_workbook
            fx_rate, fx_source = _resolve_fx_rate(
                workbook_path=fx_source_book,
                report_currency=report_currency,
                quote_currency=quote_currency,
                manual_fx=None,
            )

            rows, summary = build_sheet_rows(
                income_df_raw=income_df,
                balance_df_raw=balance_df,
                cash_df_raw=cash_df,
                info_df=info_df,
                code=code,
                years=years,
                fx_rate=fx_rate,
                yahoo_timeout=args.yahoo_timeout,
                shares_fallback_workbook=fx_workbook,
            )

            lower_series = next((x for x in rows if x.label == "Lower Price Limit"), None)
            upper_series = next((x for x in rows if x.label == "Upper Price Limit"), None)
            if lower_series is None or upper_series is None or len(lower_series.values) == 0 or len(upper_series.values) == 0:
                raise ValueError("lower/upper price rows missing")

            latest_lower_quote = float(lower_series.values[-1]) * float(fx_rate)
            latest_upper_quote = float(upper_series.values[-1]) * float(fx_rate)
            latest_close, price_source = resolve_latest_close(
                code,
                info_df,
                quote_cache=quote_cache,
                quote_source_cache=quote_source_cache,
            )

            write_sheet(
                workbook_path=workbook_path,
                sheet_name=calc_sheet_name(code),
                company=company or code,
                years=years,
                rows=rows,
                summary=summary,
                overwrite=True,
                report_currency=report_currency,
                quote_currency=quote_currency,
                fx_rate=fx_rate,
                fx_source=fx_source,
            )

            results[row] = (
                {
                    "close": round(latest_close, 2) if latest_close is not None else None,
                    "lower": round(latest_lower_quote, 2),
                    "upper": round(latest_upper_quote, 2),
                    "price_source": price_source,
                    "revenue_cagr_3y": calc_revenue_cagr_3y(rows),
                    "dividend_yield": (
                        dividend_yield_cache.get(code)
                        if code in dividend_yield_cache
                        else _dividend_yield_from_info(info_df)
                    ),
                    "price_currency": quote_currency,
                    "report_period": latest_period or "N/A",
                }
            )
            success += 1
            print(
                f"[{code}] close={(latest_close if latest_close is not None else float('nan')):.2f} "
                f"lower={latest_lower_quote:.2f} upper={latest_upper_quote:.2f} "
                f"| px_src={price_source} ccy={quote_currency} "
                f"report_period={latest_period or 'N/A'} fx={fx_rate:.6f}"
            )
        except Exception as exc:
            failed += 1
            results[row] = {
                "close": None,
                "lower": None,
                "upper": None,
                "price_source": f"error: {exc}",
                "revenue_cagr_3y": None,
                "dividend_yield": None,
                "price_currency": "",
                "report_period": "N/A",
            }
            print(f"[{code}] failed: {exc}")

    wb = load_workbook(workbook_path, data_only=False)
    ws = wb[home_sheet_name]
    ws.cell(1, 4).value = "最近收盘价"
    ws.cell(1, 5).value = "代码估值下限"
    ws.cell(1, 6).value = "代码估值上限"
    ws.cell(1, 7).value = "3年平均年营收增长率"
    ws.cell(1, 8).value = "股息率"
    ws.cell(1, 9).value = "股价币种"
    ws.cell(1, 10).value = "最新财报期"

    for row, vals in results.items():
        close_px = vals.get("close")
        low = vals.get("lower")
        high = vals.get("upper")
        rev_cagr_3y = vals.get("revenue_cagr_3y")
        dividend_yield = vals.get("dividend_yield")
        price_ccy = vals.get("price_currency")
        report_period = vals.get("report_period")
        ws.cell(row, 4).value = close_px
        ws.cell(row, 5).value = low
        ws.cell(row, 6).value = high
        ws.cell(row, 7).value = rev_cagr_3y
        ws.cell(row, 8).value = dividend_yield
        ws.cell(row, 9).value = price_ccy
        ws.cell(row, 10).value = report_period
        if close_px is not None:
            ws.cell(row, 4).number_format = "0.00"
        if low is not None:
            ws.cell(row, 5).number_format = "0.00"
        if high is not None:
            ws.cell(row, 6).number_format = "0.00"
        if rev_cagr_3y is not None:
            ws.cell(row, 7).number_format = "0.00%"
        if dividend_yield is not None:
            ws.cell(row, 8).number_format = "0.00%"

        if close_px is None or low is None or high is None:
            font_color = None
        elif close_px < low:
            font_color = ROW_FONT_RED
        elif close_px > high:
            font_color = ROW_FONT_GREEN
        else:
            font_color = ROW_FONT_GRAY

        max_col = max(10, ws.max_column)
        for c in range(1, max_col + 1):
            cell = ws.cell(row, c)
            base_font = copy(cell.font)
            cell.font = Font(
                name=base_font.name,
                sz=base_font.sz,
                b=base_font.b,
                i=base_font.i,
                charset=base_font.charset,
                u=base_font.u,
                strike=base_font.strike,
                color=(Color(rgb=f"00{font_color}") if font_color is not None else None),
                vertAlign=base_font.vertAlign,
                outline=base_font.outline,
                shadow=base_font.shadow,
                condense=base_font.condense,
                extend=base_font.extend,
                family=base_font.family,
                scheme=base_font.scheme,
            )
            cell.fill = NO_FILL

        code = row_to_code.get(row, "")
        sheet_name = calc_sheet_name(code) if code else ""
        company_cell = ws.cell(row, 2)
        if sheet_name and sheet_name in wb.sheetnames:
            company_cell.hyperlink = f"#'{sheet_name}'!A1"
            base_font = copy(company_cell.font)
            company_cell.font = Font(
                name=base_font.name,
                sz=base_font.sz,
                b=base_font.b,
                i=base_font.i,
                charset=base_font.charset,
                u="single",
                strike=base_font.strike,
                color=base_font.color,
                vertAlign=base_font.vertAlign,
                outline=base_font.outline,
                shadow=base_font.shadow,
                condense=base_font.condense,
                extend=base_font.extend,
                family=base_font.family,
                scheme=base_font.scheme,
            )
        else:
            company_cell.hyperlink = None

    wb.save(workbook_path)
    print(f"\nDone. success={success}, failed={failed}, workbook={workbook_path}")


if __name__ == "__main__":
    main()
