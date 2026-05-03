"""Parsers for local financial files."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from app.utils.financial_mapping import finalize_financial_row, map_raw_metric, parse_report_date


def parse_financial_file(path: Path) -> dict:
    """Parse a financial file into raw yearly metric records."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".json":
        return _parse_json(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    return {"kind": "unsupported", "records": [], "warnings": [f"Unsupported file type ignored: {path.name}"]}


def _parse_csv(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return {"kind": "empty", "records": [], "warnings": [f"Empty CSV ignored: {path.name}"]}
    if len(rows[0]) == 2 and rows[0][0]:
        return _parse_key_value_csv(rows)
    return _parse_statement_rows(path, rows)


def _parse_key_value_csv(rows: list[list[str]]) -> dict:
    payload: dict[str, str] = {}
    for row in rows:
        if len(row) >= 2 and row[0].strip():
            payload[row[0].strip()] = row[1].strip()
    return {"kind": "info", "records": [payload], "warnings": []}


def _parse_statement_rows(path: Path, rows: list[list[str]]) -> dict:
    headers = rows[0][1:]
    yearly: dict[int, dict] = {}
    warnings: list[str] = []
    from app.utils.financial_mapping import parse_numeric_value

    for index, header in enumerate(headers, start=1):
        year, report_date = parse_report_date(header)
        if year is None:
            continue
        yearly[year] = {"fiscal_year": year, "report_date": report_date.isoformat() if report_date else None}
        for row in rows[1:]:
            if not row or not row[0].strip():
                continue
            field = map_raw_metric(row[0].strip())
            if field is None or index >= len(row):
                continue
            value = parse_numeric_value(row[index])
            if value is not None:
                yearly[year][field] = value
        yearly[year]["raw_source"] = {"file_name": path.name}

    records = [finalize_financial_row(row) for _, row in sorted(yearly.items(), reverse=True)]
    if not records:
        warnings.append(f"No recognized metrics parsed from {path.name}")
    return {"kind": "statement", "records": records, "warnings": warnings}


def _parse_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    records: list[dict] = []
    warnings: list[str] = []
    if isinstance(payload, list):
        for item in payload:
            if isinstance(item, dict) and item.get("fiscal_year") is not None:
                records.append(finalize_financial_row(dict(item)))
    elif isinstance(payload, dict) and isinstance(payload.get("records"), list):
        for item in payload["records"]:
            if isinstance(item, dict):
                records.append(finalize_financial_row(dict(item)))
    else:
        warnings.append(f"JSON format not recognized for {path.name}")
    return {"kind": "json", "records": records, "warnings": warnings}


def _parse_xlsx(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"kind": "xlsx", "records": [], "warnings": [f"openpyxl not installed; skipped {path.name}"]}

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = [[("" if cell is None else str(cell)) for cell in row] for row in sheet.iter_rows(values_only=True)]
    return _parse_statement_rows(path, rows)
