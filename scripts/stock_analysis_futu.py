#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
from zipfile import BadZipFile

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook

try:
    from .download_financial_data import (  # type: ignore[attr-defined]  # noqa: F401
        CAGR,
        downloadFinancialDataFutu,
        getStockInfoWeb,
        getValue,
        growthRateCalculation,
        movingAverage,
        strToNum,
    )
except Exception:
    try:
        from download_financial_data import (  # noqa: F401
            CAGR,
            downloadFinancialDataFutu,
            getStockInfoWeb,
            getValue,
            growthRateCalculation,
            movingAverage,
            strToNum,
        )
    except Exception:
        def strToNum(ls):
            num = np.zeros(len(ls))
            for i in range(len(ls)):
                text = str(ls.iloc[i]).replace(",", "")
                if re.search(r"\d", str(ls.iloc[i])) is None:
                    num[i] = 0
                elif "萬億" in str(ls.iloc[i]):
                    num[i] = float(re.findall(r"-?\d+(?:\.\d+)?", text)[0]) * 1000000000000
                elif "億" in str(ls.iloc[i]):
                    num[i] = float(re.findall(r"-?\d+(?:\.\d+)?", text)[0]) * 100000000
                elif "萬" in str(ls.iloc[i]):
                    num[i] = float(re.findall(r"-?\d+(?:\.\d+)?", text)[0]) * 10000
                elif "%" in str(ls.iloc[i]):
                    num[i] = float(re.findall(r"-?\d+(?:\.\d+)?", text)[0]) * 0.01
                else:
                    num[i] = float(text)
            return num

        def getValue(data, parameters, item):
            ls = parameters.loc[item].dropna()
            result = np.zeros(len(data.columns))
            for i in range(len(ls)):
                try:
                    value = strToNum(data.loc[ls.iloc[i]])
                except Exception:
                    value = np.zeros(len(data.columns))
                result = value if i == 0 else (result + value)
            return result

        def movingAverage(ls):
            average = ls.copy()
            for i in range(len(ls) - 1):
                average[i] = (ls[i] + ls[i + 1]) / 2
            return average

        def growthRateCalculation(ls):
            rate = np.zeros(len(ls))
            for i in range(len(ls) - 1):
                rate[i] = (ls[i] - ls[i + 1]) / ls[i + 1]
                if abs(rate[i]) > 2:
                    rate[i] = 0
            return rate

        def CAGR(ls):
            return pow(ls[0] / ls[len(ls) - 1], 1 / (len(ls) - 1)) - 1

        def getStockInfoWeb(code):
            raise RuntimeError(f"getStockInfoWeb unavailable for {code}: downloader dependencies not installed")

        def downloadFinancialDataFutu(code, yearCycle=True):
            raise RuntimeError(f"downloadFinancialDataFutu unavailable for {code}: downloader dependencies not installed")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Futu financial analysis runner (input stock codes directly)"
    )
    parser.add_argument(
        "codes",
        nargs="+",
        help=(
            "Stock code(s), e.g. SH.600519 HK.00700 US.AAPL. "
            "Can also pass raw code like 600519, 000001, 00700, AAPL."
        ),
    )
    parser.add_argument(
        "--download",
        action="store_true",
        help="Download stock info + financial statements before analysis",
    )
    parser.add_argument(
        "--plot",
        action="store_true",
        help="Render and save plots; disabled by default so the script only computes metrics",
    )
    parser.add_argument(
        "--summary-file",
        default="results/analysis_summary.xlsx",
        help="Path to save summary excel",
    )
    parser.add_argument(
        "--sheet",
        default="summary",
        help="Sheet name in summary excel",
    )
    parser.add_argument(
        "--min-years",
        type=int,
        default=3,
        help="Minimum required common annual periods across balance/income/cash",
    )
    parser.add_argument(
        "--retry",
        type=int,
        default=3,
        help="Retry times for download when integrity check fails",
    )
    parser.add_argument(
        "--max-years",
        type=int,
        default=10,
        help="Maximum number of recent annual periods used for analysis/plot",
    )
    parser.add_argument(
        "--value-file",
        default="Value_Investment.xlsx",
        help="Workbook used for code->name mapping (sheet from --stats-sheet)",
    )
    parser.add_argument(
        "--stats-sheet",
        default="统计",
        help="Sheet name in --value-file for company names/codes",
    )
    parser.add_argument(
        "--quote-timeout",
        type=int,
        default=12,
        help="Realtime quote HTTP timeout in seconds",
    )
    parser.add_argument(
        "--skip-live-quote-refresh",
        action="store_true",
        help="Use existing info csv for quote fields without refreshing from web",
    )
    return parser.parse_args()


def normalize_code(raw_code: str) -> str:
    text = raw_code.strip().upper()
    if not text:
        return text
    if "." in text:
        return text

    if re.fullmatch(r"\d{6}", text):
        if text.startswith(("6", "9")):
            return f"SH.{text}"
        if text.startswith(("0", "3")):
            return f"SZ.{text}"
    if re.fullmatch(r"\d{5}", text):
        return f"HK.{text}"
    if re.fullmatch(r"\d{4}", text):
        return f"HK.{text.zfill(5)}"
    return f"US.{text}"


def parse_code_inputs(codes: Sequence[str]) -> List[str]:
    result: List[str] = []
    seen = set()
    for code in codes:
        parts = [x.strip() for x in code.split(",") if x.strip()]
        for part in parts:
            normalized = normalize_code(part)
            if normalized and normalized not in seen:
                seen.add(normalized)
                result.append(normalized)
    return result


def _safe_float(v: object) -> Optional[float]:
    if v is None:
        return None
    try:
        f = float(v)
    except Exception:
        return None
    if pd.isna(f):
        return None
    return f


def _info_cell_to_text(value: object) -> str:
    if isinstance(value, pd.DataFrame):
        if value.empty:
            return ""
        value = value.iloc[0, 0]
    elif isinstance(value, pd.Series):
        if value.empty:
            return ""
        value = value.iloc[0]
    return str(value).strip()


def _parse_number_text(value: object, with_units: bool = False) -> Optional[float]:
    text = _info_cell_to_text(value).replace(",", "").replace(" ", "")
    if not text or text in {"--", "nan", "None"}:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", text)
    if not m:
        return None
    try:
        num = float(m.group(0))
    except Exception:
        return None
    if not with_units:
        return num
    if "萬億" in text or "万亿" in text:
        num *= 1_0000_0000_0000
    elif "億" in text or "亿" in text:
        num *= 1_0000_0000
    elif "萬" in text or "万" in text:
        num *= 1_0000
    elif "%" in text:
        num *= 0.01
    return num


def _read_info_numeric(info_data: pd.DataFrame, keys: Sequence[str], with_units: bool = False) -> Optional[float]:
    if info_data is None or info_data.empty:
        return None
    index_text = [str(x).strip() for x in info_data.index]

    for key in keys:
        for idx in index_text:
            if idx == key:
                try:
                    parsed = _parse_number_text(info_data.loc[idx], with_units=with_units)
                except Exception:
                    continue
                if parsed is not None:
                    return parsed

    for key in keys:
        for idx in index_text:
            if key and key in idx:
                try:
                    parsed = _parse_number_text(info_data.loc[idx], with_units=with_units)
                except Exception:
                    continue
                if parsed is not None:
                    return parsed
    return None


def _code_to_yahoo_symbol(code: str) -> Optional[str]:
    c = normalize_code(code)
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "HK":
        return f"{int(ticker):04d}.HK" if ticker.isdigit() else f"{ticker}.HK"
    if market == "SH":
        return f"{ticker}.SS"
    if market == "SZ":
        return f"{ticker}.SZ"
    if market == "US":
        return ticker
    if market == "JP":
        return f"{ticker}.T"
    return None


def _code_to_tencent_symbol(code: str) -> Optional[str]:
    c = normalize_code(code)
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "HK":
        return f"hk{ticker}"
    if market == "SH":
        return f"sh{ticker}"
    if market == "SZ":
        return f"sz{ticker}"
    if market == "US":
        return f"us{ticker}"
    if market == "JP":
        return f"jp{ticker}"
    return None


def load_name_map_from_value_file(project_root: Path, value_file: str, sheet_name: str) -> Dict[str, str]:
    path = (project_root / value_file).resolve()
    if not path.exists():
        return {}
    try:
        df = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
    except Exception as exc:
        print(f"read value file failed ({path}#{sheet_name}): {exc}")
        return {}
    if df.empty or "代码" not in df.columns:
        return {}

    name_col = "公司" if "公司" in df.columns else ("名称" if "名称" in df.columns else "")
    if not name_col:
        return {}

    out: Dict[str, str] = {}
    for _, row in df.iterrows():
        code = normalize_code(str(row.get("代码", "")).strip().upper())
        name = str(row.get(name_col, "")).strip()
        if code and name and code not in out:
            out[code] = name
    return out


def fetch_live_quote_metrics(codes: List[str], timeout: int = 12, refresh: bool = True) -> Dict[str, Dict[str, object]]:
    codes = [normalize_code(c) for c in codes if str(c).strip()]
    codes = list(dict.fromkeys(codes))
    out: Dict[str, Dict[str, object]] = {c: {} for c in codes}
    if not codes:
        return out

    # 1) Load existing local info metrics first.
    def _read_info(code: str) -> Dict[str, object]:
        info_path = Path(__file__).resolve().parents[1] / "data" / "finance_report_web" / f"{code}_info.csv"
        if not info_path.exists():
            return {}
        try:
            info = pd.read_csv(info_path, header=None, index_col=0, encoding="utf_8_sig")
        except Exception:
            return {}
        bucket: Dict[str, object] = {}
        try:
            if "股票名称" in info.index:
                bucket["name"] = _info_cell_to_text(info.loc["股票名称", 1]).strip()
        except Exception:
            pass
        price = _read_info_numeric(info, ["平均价", "最新价", "现价"], with_units=False)
        if price is not None and price > 0:
            bucket["price"] = float(price)
        pe = _read_info_numeric(info, ["市盈率TTM", "市盈率(静)", "市盈率(靜)", "市盈率", "PE"], with_units=False)
        if pe is not None:
            bucket["pe"] = float(pe)
        pb = _read_info_numeric(info, ["市净率TTM", "市淨率TTM", "市净率", "市淨率", "PB", "P/B"], with_units=False)
        if pb is not None:
            bucket["pb"] = float(pb)
        mv = _read_info_numeric(info, ["总市值", "總市值", "市值"], with_units=True)
        if mv is not None and mv > 0:
            bucket["market_cap"] = float(mv)
        dv = _read_info_numeric(info, ["股息率", "股息率TTM", "股息"], with_units=True)
        if dv is not None:
            bucket["div_yield"] = float(dv)
        return bucket

    # First read local info as fallback snapshot.
    for c in codes:
        m = _read_info(c)
        if m.get("price", 0):
            out[c] = m

    # 2) Refresh all codes from Futu webpage via existing selenium scraper.
    if not refresh:
        got_price = sum(1 for c in codes if "price" in out.get(c, {}))
        print(f"Futu live quotes fetched: {got_price}/{len(codes)} (refreshed 0)")
        return out

    refreshed = 0
    for c in codes:
        try:
            getStockInfoWeb(c)
            refreshed += 1
        except Exception as exc:
            print(f"futu live quote refresh failed {c}: {exc}")
            continue
        m = _read_info(c)
        if m.get("price", 0):
            out[c] = m

    got_price = sum(1 for c in codes if "price" in out.get(c, {}))
    print(f"Futu live quotes fetched: {got_price}/{len(codes)} (refreshed {refreshed})")
    return out


def safe_div(numerator: np.ndarray, denominator: np.ndarray, ndigits: int = 3) -> np.ndarray:
    result = np.divide(
        numerator,
        denominator,
        out=np.zeros_like(numerator, dtype=float),
        where=np.abs(denominator) > 1e-12,
    )
    return np.round(result, ndigits)


def normalize_financial_cell(value: object) -> str:
    """Normalize mixed strings like '2,027.18%239.67萬' into amount-only text."""
    if pd.isna(value):
        return "--"

    text = str(value).strip().replace(" ", "")
    if text in {"", "nan", "None"}:
        return "--"

    # Keep amount part after YoY percentage.
    if "%" in text:
        tail = text.rsplit("%", 1)[-1].strip()
        if re.search(r"\d", tail):
            text = tail

    # Convert malformed leading '--123.4萬' / '---123.4萬' to '-123.4萬'.
    text = re.sub(r"^--+(?=\d)", "-", text)
    return text


def sanitize_financial_frame(frame: pd.DataFrame) -> pd.DataFrame:
    clean = frame.copy()
    for col in clean.columns:
        clean[col] = clean[col].map(normalize_financial_cell)
    return clean


def safe_mean(arr: np.ndarray, n: int = 3) -> float:
    sub = np.array(arr[:n], dtype=float)
    sub = sub[np.isfinite(sub)]
    if sub.size == 0:
        return 0.0
    return float(np.mean(sub))


def ensure_columns_aligned(
    balance_data: pd.DataFrame,
    income_data: pd.DataFrame,
    cash_data: pd.DataFrame,
    max_years: int = 10,
) -> Tuple[List[str], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    d1 = [str(c) for c in balance_data.columns]
    d2 = [str(c) for c in income_data.columns]
    d3 = [str(c) for c in cash_data.columns]
    common_dates = [d for d in d1 if d in d2 and d in d3]

    if not common_dates:
        raise ValueError("No common report periods across balance/income/cash data")

    # Keep only the most recent N annual periods for faster plotting and cleaner charts.
    if max_years > 0 and len(common_dates) > max_years:
        def _sort_key(period: str):
            text = str(period).strip()
            try:
                return pd.to_datetime(text, format="%Y/%m/%d", errors="raise")
            except Exception:
                return pd.to_datetime("1900-01-01")

        sorted_desc = sorted(common_dates, key=_sort_key, reverse=True)
        common_dates = sorted_desc[:max_years]

    return (
        common_dates,
        balance_data[common_dates],
        income_data[common_dates],
        cash_data[common_dates],
    )


def validate_statement_integrity(
    balance_data: pd.DataFrame,
    income_data: pd.DataFrame,
    cash_data: pd.DataFrame,
    min_years: int = 3,
) -> Tuple[bool, str]:
    d_cols = [str(c) for c in balance_data.columns]
    b_cols = [str(c) for c in income_data.columns]
    c_cols = [str(c) for c in cash_data.columns]
    common = [x for x in d_cols if x in b_cols and x in c_cols]

    counts = [len(d_cols), len(b_cols), len(c_cols)]
    reasons: List[str] = []
    if min(counts) < min_years:
        reasons.append(
            f"at least one statement has <{min_years} periods (balance={len(d_cols)}, income={len(b_cols)}, cash={len(c_cols)})"
        )
    if len(common) < min_years:
        reasons.append(f"common periods across three statements are only {len(common)}")

    if reasons:
        return False, "; ".join(reasons)
    return True, ""


def load_data(project_root: Path, code: str) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    parameter_root = Path(
        os.getenv(
            "INVESTMENT_MAS_LEGACY_FINANCIAL_PARAMETER_DIR",
            str(project_root / "data" / "parameters"),
        )
    )
    parameters_path = parameter_root / f"{code[:2]}_parameters.csv"
    finance_root = project_root / "data" / "finance_report_web"

    info_path = finance_root / f"{code}_info.csv"
    balance_path = finance_root / f"{code}_balance_year.csv"
    income_path = finance_root / f"{code}_income_year.csv"
    cash_path = finance_root / f"{code}_cash_year.csv"

    required = [parameters_path, balance_path, income_path, cash_path]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError(f"Missing files for {code}: {missing}")

    parameters = pd.read_csv(parameters_path, index_col=0, encoding="utf-8-sig")
    if info_path.exists():
        info_data = sanitize_financial_frame(pd.read_csv(info_path, index_col=0, encoding="utf-8-sig"))
    else:
        info_data = pd.DataFrame()
    balance_data = sanitize_financial_frame(pd.read_csv(balance_path, index_col=0, encoding="utf-8-sig"))
    income_data = sanitize_financial_frame(pd.read_csv(income_path, index_col=0, encoding="utf-8-sig"))
    cash_data = sanitize_financial_frame(pd.read_csv(cash_path, index_col=0, encoding="utf-8-sig"))

    return parameters, info_data, balance_data, income_data, cash_data


def verify_code_integrity(project_root: Path, code: str, min_years: int = 3) -> Tuple[bool, str]:
    _, _, balance_data, income_data, cash_data = load_data(project_root, code)
    return validate_statement_integrity(balance_data, income_data, cash_data, min_years=min_years)


def read_info_metric(info_data: pd.DataFrame, key: str, default: float = 0.0) -> float:
    if info_data is None or info_data.empty:
        return default
    if key not in info_data.index:
        return default
    try:
        value = float(strToNum(info_data.loc[key])[0])
        if pd.isna(value):
            return default
        return value
    except Exception:
        return default


def read_info_metric_any(info_data: pd.DataFrame, keys: Sequence[str], default: float = 0.0) -> float:
    value = _read_info_numeric(info_data, keys, with_units=True)
    if value is None:
        return default
    return float(value)


def save_report(report_path: Path, tables: List[pd.DataFrame]) -> None:
    for i, table in enumerate(tables):
        mode = "w" if i == 0 else "a"
        table.T.to_csv(
            report_path,
            mode=mode,
            encoding="utf_8_sig",
            header=False,
            index=True,
        )


def annotate_all(ax, date_count: int, frame: pd.DataFrame, fmt: str) -> None:
    for row in frame.itertuples():
        for value in row[2:]:
            ax.text(date_count - 1 - row.Index, value, format(value, fmt))


def _annual_metric_records(date: List[str], earning_power: pd.DataFrame) -> List[Dict[str, object]]:
    """Convert earning-power table into annual metric rows."""

    records: List[Dict[str, object]] = []
    for idx, period in enumerate(date):
        row = earning_power.iloc[idx]
        try:
            fiscal_year = int(str(period).split("/")[0])
        except Exception:
            fiscal_year = 0
        records.append(
            {
                "fiscal_year": fiscal_year,
                "report_date": str(period),
                "gross_margin": _safe_float(row.get("grossProfitMargin")),
                "net_margin": _safe_float(row.get("netProfitMargin")),
                "roe": _safe_float(row.get("ROE")),
                "roa": _safe_float(row.get("ROA")),
                "roic": _safe_float(row.get("ROIC")),
            }
        )
    return records


def plot_report(
    code: str,
    stock_name: str,
    last_price: float,
    pb_value: float,
    pe_value: float,
    dv_value: float,
    cash_return_rate: float,
    date: List[str],
    tables: Dict[str, pd.DataFrame],
    output_png: Path,
    show_plot: bool,
) -> None:
    import matplotlib.pyplot as plt

    asset_plot = tables["asset"].copy()
    balance_plot = tables["balance"].copy()
    for frame in [asset_plot, balance_plot]:
        num_cols = [c for c in frame.columns if c != "Date"]
        frame[num_cols] = frame[num_cols].clip(lower=0)

    earning_plot = tables["earning_power"].copy()
    expense_plot = tables["expense_rate"].copy()
    growth_plot = tables["growth_rate"].copy()
    safe_plot = tables["safe_factor"].copy()
    op_plot = tables["operating_capability"].copy()
    for frame, lo, hi in [
        (earning_plot, -1.0, 1.5),
        (expense_plot, -1.0, 0.2),
        (growth_plot, -2.0, 2.0),
        (safe_plot, 0.1, 30.0),
        (op_plot, 0.0, 10.0),
    ]:
        num_cols = [c for c in frame.columns if c != "Date"]
        frame[num_cols] = frame[num_cols].clip(lower=lo, upper=hi)

    fig, axes = plt.subplots(7, 2, figsize=(30, 40))
    fig.suptitle(
        f"{code} Financial Report\n"
        f"Cash Return Rate: {cash_return_rate:.2%}\n"
        f"Price: {last_price}\n"
        f"PB(A): {pb_value}  PE: {pe_value}  DV: {dv_value}",
        fontsize=20,
        fontweight="bold",
    )

    limit = pd.DataFrame({"Date": date, "Limit": np.ones(len(date))})
    limit.plot.line(ax=axes[3, 1], x="Date", linestyle="--", color="r", rot=0)
    limit.plot.line(ax=axes[6, 1], x="Date", linestyle="--", color="r", rot=0)

    asset_plot.iloc[::-1].plot.area(ax=axes[0, 0], x="Date", title="Assets", rot=0)
    balance_plot.iloc[::-1].plot.area(ax=axes[1, 0], x="Date", title="Balance", rot=0)
    tables["income"].iloc[::-1].plot.bar(ax=axes[2, 0], x="Date", title="Income", rot=0)
    tables["cash"].iloc[::-1].plot.bar(ax=axes[3, 0], x="Date", title="CashFlow", rot=0)
    tables["operating_cash"].iloc[::-1].plot.bar(ax=axes[4, 0], x="Date", title="Operating Cash", rot=0)
    tables["investment_cash"].iloc[::-1].plot.bar(ax=axes[5, 0], x="Date", title="Investment Cash", rot=0)
    tables["financial_cash"].iloc[::-1].plot.bar(ax=axes[6, 0], x="Date", title="Financial Cash", rot=0)

    tables["index_per_share"].iloc[::-1].plot.bar(ax=axes[0, 1], x="Date", title="Index Per Share", rot=0)
    earning_plot.iloc[::-1].plot.line(ax=axes[1, 1], x="Date", marker="o", title="Earning Power", rot=0)
    growth_plot.iloc[::-1].plot.bar(ax=axes[2, 1], x="Date", title="Growth Rate", rot=0)
    safe_plot.iloc[::-1].plot.bar(ax=axes[3, 1], x="Date", title="Safe Factor", logy=True, rot=0)
    expense_plot.iloc[::-1].plot.line(ax=axes[4, 1], x="Date", marker="o", title="Expense Rate", rot=0)
    tables["operating_funds"].iloc[::-1].plot.bar(ax=axes[5, 1], x="Date", title="Operating Funds", rot=0)
    op_plot.iloc[::-1].plot.line(ax=axes[6, 1], x="Date", title="Operation", logy=True, rot=0)

    annotate_all(axes[0, 1], len(date), tables["index_per_share"], ".1f")
    annotate_all(axes[1, 1], len(date), earning_plot, ".0%")
    annotate_all(axes[2, 1], len(date), growth_plot, ".0%")
    annotate_all(axes[3, 1], len(date), safe_plot, ".1f")
    annotate_all(axes[4, 1], len(date), expense_plot, ".0%")
    annotate_all(axes[6, 1], len(date), op_plot, ".1f")

    for ax in axes.ravel():
        ax.legend(loc="upper left", fontsize=10)

    fig.savefig(output_png)
    if show_plot:
        plt.show()
    plt.close(fig)


def financial_analysis(
    project_root: Path,
    code: str,
    show_plot: bool = False,
    max_years: int = 10,
    name_map: Optional[Dict[str, str]] = None,
    live_quotes: Optional[Dict[str, Dict[str, object]]] = None,
    write_artifacts: bool = False,
    return_details: bool = False,
) -> pd.DataFrame | Tuple[pd.DataFrame, Dict[str, object]]:
    parameters, info_data, balance_data, income_data, cash_data = load_data(project_root, code)
    ok, reason = validate_statement_integrity(balance_data, income_data, cash_data, min_years=3)
    if not ok:
        raise ValueError(f"data integrity check failed for {code}: {reason}")
    date, balance_data, income_data, cash_data = ensure_columns_aligned(
        balance_data,
        income_data,
        cash_data,
        max_years=max_years,
    )

    code = normalize_code(code)
    quote = (live_quotes or {}).get(code, {})

    info_stock_name = ""
    if info_data is not None and (not info_data.empty) and ("股票名称" in info_data.index):
        try:
            info_stock_name = str(info_data.loc["股票名称"].iloc[0]).strip()
        except Exception:
            info_stock_name = ""

    stock_name = (
        (name_map or {}).get(code)
        or info_stock_name
        or str(quote.get("name", "")).strip()
        or code
    )

    last_price = read_info_metric_any(info_data, ["平均价", "最新价", "现价"], default=0.0)
    market_value = read_info_metric_any(info_data, ["总市值", "總市值", "市值"], default=0.0)
    total_shares = read_info_metric_any(info_data, ["总股本", "總股本"], default=0.0)
    pe_value = read_info_metric_any(info_data, ["市盈率TTM", "市盈率(静)", "市盈率(靜)", "市盈率", "PE"], default=0.0)
    pb_value = read_info_metric_any(info_data, ["市净率TTM", "市淨率TTM", "市净率", "市淨率", "PB", "P/B"], default=0.0)
    dv_value = round(read_info_metric_any(info_data, ["股息率", "股息率TTM", "股息"], default=0.0), 3)

    if _safe_float(quote.get("price")) is not None:
        last_price = float(quote["price"])
    if _safe_float(quote.get("pe")) is not None:
        pe_value = float(quote["pe"])
    if _safe_float(quote.get("pb")) is not None:
        pb_value = float(quote["pb"])
    if _safe_float(quote.get("market_cap")) is not None:
        market_value = float(quote["market_cap"])
    if _safe_float(quote.get("div_yield")) is not None:
        dv_value = round(float(quote["div_yield"]), 3)
    if total_shares <= 0 and last_price > 0:
        total_shares = market_value / last_price

    cash = getValue(balance_data, parameters, "cash")
    accounts_receivable = getValue(balance_data, parameters, "accountsReceivable")
    inventory = getValue(balance_data, parameters, "inventory")
    advance_payment = getValue(balance_data, parameters, "advancePayment")
    restricted_cash = getValue(balance_data, parameters, "restrictedCash")
    financial_assets = getValue(balance_data, parameters, "financialAssets")
    current_assets = getValue(balance_data, parameters, "currentAssets")
    other_current_assets = (
        current_assets
        - cash
        - inventory
        - accounts_receivable
        - restricted_cash
        - financial_assets
        - advance_payment
    )

    fixed_assets = getValue(balance_data, parameters, "fixedAssets")
    good_will = getValue(balance_data, parameters, "goodWill")
    intangible_assets = getValue(balance_data, parameters, "intangibleAssets")
    long_term_accounts_receivable = getValue(balance_data, parameters, "longTermAccountsReceivable")
    long_term_advance_payment = getValue(balance_data, parameters, "longTermAdvancePayment")
    long_term_investment = getValue(balance_data, parameters, "longTermInvestment")
    deferred_assets = getValue(balance_data, parameters, "deferredAssets")
    non_current_financial_assets = getValue(balance_data, parameters, "nonCurrentFinancialAssets")
    non_current_assets = getValue(balance_data, parameters, "nonCurrentAssets")
    other_non_current_assets = (
        non_current_assets
        - fixed_assets
        - good_will
        - intangible_assets
        - long_term_accounts_receivable
        - long_term_advance_payment
        - long_term_investment
        - deferred_assets
        - non_current_financial_assets
    )
    total_assets = getValue(balance_data, parameters, "totalAssets")

    accounts_payable = getValue(balance_data, parameters, "accountsPayable")
    short_term_loan = getValue(balance_data, parameters, "shortTermLoan")
    deferred_revenue = getValue(balance_data, parameters, "deferredRevenue")
    advance_receipt = getValue(balance_data, parameters, "advanceReceipt")
    current_liability = getValue(balance_data, parameters, "currentLiability")
    other_current_liability = (
        current_liability - accounts_payable - short_term_loan - deferred_revenue - advance_receipt
    )

    long_term_loan = getValue(balance_data, parameters, "longTermLoan")
    long_term_lease_liabilities = getValue(balance_data, parameters, "longTermLeaseLiabilities")
    long_term_payable = getValue(balance_data, parameters, "longTermPayable")
    deferred_liability = getValue(balance_data, parameters, "deferredLiability")
    non_current_liability = getValue(balance_data, parameters, "nonCurrentLiability")
    other_non_current_liability = (
        non_current_liability
        - long_term_loan
        - long_term_lease_liabilities
        - long_term_payable
        - deferred_liability
    )
    total_debts = getValue(balance_data, parameters, "totalDebts")

    stockholders_equity = getValue(balance_data, parameters, "stockholdersEquity")
    non_control_interest = getValue(balance_data, parameters, "nonControlInterest")

    operating_income = getValue(income_data, parameters, "operatingIncome")
    operating_cost = -np.abs(getValue(income_data, parameters, "operatingCost"))
    marketing_expense = -np.abs(getValue(income_data, parameters, "marketingExpense"))
    administrative_expense = -np.abs(getValue(income_data, parameters, "administrativeExpense"))
    rd_expense = -np.abs(getValue(income_data, parameters, "RDExpense"))
    financial_cost = getValue(income_data, parameters, "financialCost")
    devaluation = getValue(income_data, parameters, "devaluation")
    operating_profit = getValue(income_data, parameters, "operatingProfit")
    pre_tax_profit = getValue(income_data, parameters, "preTaxProfit")
    income_tax = -np.abs(getValue(income_data, parameters, "incomeTax"))
    net_profit = getValue(income_data, parameters, "netProfit")

    eps = getValue(income_data, parameters, "EPS")
    deps = getValue(income_data, parameters, "DEPS")
    dps = getValue(income_data, parameters, "DPS")

    depreciation = getValue(cash_data, parameters, "depreciation")
    accounts_receivable_change = getValue(cash_data, parameters, "accountsReceivableChange")
    accounts_payable_change = getValue(cash_data, parameters, "accountsPayableChange")
    advanced_receivable_change = getValue(cash_data, parameters, "advancedReceivableChange")
    advanced_payable_change = getValue(cash_data, parameters, "advancedPayableChange")
    inventory_change = getValue(cash_data, parameters, "inventoryChange")
    tax_paid = -np.abs(getValue(cash_data, parameters, "taxPaid"))
    operating_net_cash = getValue(cash_data, parameters, "operatingNetCash")

    operation_adjust = (
        operating_net_cash
        - depreciation
        - accounts_receivable_change
        - accounts_payable_change
        - advanced_receivable_change
        - advanced_payable_change
        - inventory_change
        - tax_paid
        - pre_tax_profit
    )

    property_sale = getValue(cash_data, parameters, "propertySale")
    property_purchase = -np.abs(getValue(cash_data, parameters, "propertyPurchase"))
    intangible_asset_sale = getValue(cash_data, parameters, "intangibleAssetSale")
    intangible_asset_purchase = -np.abs(getValue(cash_data, parameters, "intangibleAssetPurchase"))
    sub_company_sale = getValue(cash_data, parameters, "subCompanySale")
    sub_company_purchase = -np.abs(getValue(cash_data, parameters, "subCompanyPurchase"))
    investment_cash_return = getValue(cash_data, parameters, "investmentCashReturn")
    investment_cash_pay = -np.abs(getValue(cash_data, parameters, "investmentCashPay"))
    investment_net_cash = getValue(cash_data, parameters, "investmentNetCash")

    property_change = property_sale + property_purchase
    intangible_asset_change = intangible_asset_sale + intangible_asset_purchase
    sub_company_change = sub_company_sale + sub_company_purchase
    investment_cash_change = investment_cash_return + investment_cash_pay
    investment_adjust = (
        investment_net_cash
        - property_change
        - intangible_asset_change
        - sub_company_change
        - investment_cash_change
    )

    new_loan = getValue(cash_data, parameters, "newLoan")
    pay_back = -np.abs(getValue(cash_data, parameters, "payBack"))
    issued_shares = getValue(cash_data, parameters, "issuedShares")
    issued_bonds = getValue(cash_data, parameters, "issuedBonds")
    paid_interest = -np.abs(getValue(cash_data, parameters, "paidInterest"))
    paid_dividend = -np.abs(getValue(cash_data, parameters, "paidDividend"))
    investment_absorption = getValue(cash_data, parameters, "investmentAbsorption")
    financial_net_cash = getValue(cash_data, parameters, "financialNetCash")

    financial_adjust = (
        financial_net_cash
        - new_loan
        - pay_back
        - issued_shares
        - issued_bonds
        - paid_interest
        - paid_dividend
        - investment_absorption
    )

    marketing_expense_rate = safe_div(marketing_expense, operating_income)
    administrative_expense_rate = safe_div(administrative_expense, operating_income)
    rd_expense_rate = safe_div(rd_expense, operating_income)
    financial_cost_rate = safe_div(financial_cost, operating_income)
    income_tax_rate = safe_div(income_tax, pre_tax_profit)

    depreciation_base = fixed_assets + inventory + accounts_receivable + good_will + intangible_assets
    depreciation_cost_rate = -safe_div(depreciation, depreciation_base)
    devaluation_cost_rate = -safe_div(devaluation, depreciation_base)

    current_ratio = safe_div(current_assets, current_liability)
    quick_ratio = safe_div(current_assets - inventory, current_liability)

    gross_profit = operating_income + operating_cost
    gross_profit_margin = safe_div(gross_profit, operating_income)
    net_profit_margin = safe_div(net_profit, operating_income)

    roe = safe_div(net_profit, movingAverage(stockholders_equity))
    roa = safe_div(net_profit, movingAverage(total_assets))

    roic_base = total_debts + stockholders_equity - accounts_payable - deferred_revenue - advance_receipt
    roic = np.round(safe_div((operating_profit - financial_cost) * (1 + income_tax_rate), roic_base), 3)

    operating_income_growth_rate = np.round(growthRateCalculation(operating_income), 3)
    net_profit_growth_rate = np.round(growthRateCalculation(net_profit), 3)
    operating_income_cagr = round(CAGR(operating_income), 3)
    net_profit_cagr = round(CAGR(net_profit), 3)

    total_assets_turn_over = safe_div(operating_income, movingAverage(total_assets))
    fixed_assets_turn_over = safe_div(operating_income, movingAverage(fixed_assets))
    inventory_turn_over = safe_div(operating_income, movingAverage(inventory))
    accounts_receivable_turn_over = safe_div(operating_income, movingAverage(accounts_receivable))
    finance_leverage_degree = safe_div(movingAverage(total_assets), movingAverage(stockholders_equity))

    free_cash_flow = operating_net_cash + investment_net_cash
    net_interest_bearing_debt = short_term_loan + long_term_loan
    total_cost = market_value + net_interest_bearing_debt[0] - cash[0] - financial_assets[0]
    cash_return_rate = round((free_cash_flow[0] / total_cost), 3) if abs(total_cost) > 1e-12 else 0.0

    estimated_cagr = min(np.average(net_profit_growth_rate[:5]), 0.3)
    if abs(total_shares) > 1e-12:
        purchase_point = (
            free_cash_flow[0] * pow((1 + estimated_cagr), 3) * 25 / 2
            + cash[0]
            + financial_assets[0]
            - total_debts[0]
        ) / total_shares
    else:
        purchase_point = 0.0

    bv = np.round(stockholders_equity - good_will - intangible_assets, 3)
    net_assets = total_assets[0] - total_debts[0]
    if pb_value <= 0 and abs(net_assets) > 1e-12:
        pb_value = market_value / net_assets
    adjusted_pb = round(pb_value * net_assets / bv[0], 3) if abs(bv[0]) > 1e-12 else 0.0

    print(f"\n{stock_name}")
    print(f"Price: {last_price}")
    print(f"Cash Return Rate: {cash_return_rate:.2%}")
    print(f"Purchase Point: {round(purchase_point, 2)}")

    asset_table = pd.DataFrame(
        {
            "Date": date,
            "cash": cash,
            "accountsReceivable": accounts_receivable,
            "inventory": inventory,
            "advancePayment": advance_payment,
            "restrictedCash": restricted_cash,
            "financialAssets": financial_assets,
            "otherCurrentAssets": np.maximum(other_current_assets, 0),
            "fixedAssets": fixed_assets,
            "goodWill": good_will,
            "intangibleAssets": intangible_assets,
            "longTermAccountsReceivable": long_term_accounts_receivable,
            "longTermAdvancePayment": long_term_advance_payment,
            "longTermInvestment": long_term_investment,
            "nonCurrentFinancialAssets": non_current_financial_assets,
            "deferredAssets": deferred_assets,
            "otherNonCurrentAssets": np.maximum(other_non_current_assets, 0),
        }
    )

    balance_table = pd.DataFrame(
        {
            "Date": date,
            "accountsPayable": accounts_payable,
            "shortTermLoan": short_term_loan,
            "deferredRevenue": deferred_revenue,
            "Advance Receipt": advance_receipt,
            "otherCurrentLiability": np.maximum(other_current_liability, 0),
            "longTermLoan": long_term_loan,
            "longTermLeaseLiabilities": np.maximum(long_term_lease_liabilities, 0),
            "longTermPayable": long_term_payable,
            "deferredLiability": np.maximum(deferred_liability, 0),
            "otherNoncurrentLiability": np.maximum(other_non_current_liability, 0),
            "stockholdersEquity": np.maximum(stockholders_equity, 0),
            "nonControlInterest": np.maximum(non_control_interest, 0),
        }
    )

    income_table = pd.DataFrame(
        {
            "Date": date,
            "grossProfit": gross_profit,
            "operatingProfit": operating_profit,
            "preTaxProfit": pre_tax_profit,
            "netProfit": net_profit,
        }
    )

    operating_cash_table = pd.DataFrame(
        {
            "Date": date,
            "preTaxProfit": pre_tax_profit,
            "depreciation": depreciation,
            "inventoryChange": inventory_change,
            "accountsReceivableChange": accounts_receivable_change,
            "accountsPayableChange": accounts_payable_change,
            "advancedReceivableChange": advanced_receivable_change,
            "advancedPayableChange": advanced_payable_change,
            "taxPaid": tax_paid,
            "operationAdjust": operation_adjust,
        }
    )

    investment_cash_table = pd.DataFrame(
        {
            "Date": date,
            "propertyChange": property_change,
            "intangibleAssetChange": intangible_asset_change,
            "subCompanyChange": sub_company_change,
            "investmentCashChange": investment_cash_change,
            "investmentAdjust": investment_adjust,
        }
    )

    financial_cash_table = pd.DataFrame(
        {
            "Date": date,
            "loanChange": new_loan + pay_back,
            "issuedShares": issued_shares,
            "issuedBonds": issued_bonds,
            "paidInterest": paid_interest,
            "paidDividend": paid_dividend,
            "investmentAbsorption": investment_absorption,
            "financialAdjust": financial_adjust,
        }
    )

    operating_funds = pd.DataFrame(
        {
            "Date": date,
            "operatingAssets": current_assets - cash,
            "operatingLiability": current_liability - short_term_loan,
            "netOperatingAssets": (current_assets - cash) - (current_liability - short_term_loan),
        }
    )

    cash_table = pd.DataFrame(
        {
            "Date": date,
            "FCF": free_cash_flow,
            "netProfit": net_profit,
            "operatingNetCash": operating_net_cash,
            "investmentNetCash": investment_net_cash,
            "financialNetCash": financial_net_cash,
        }
    )

    expense_rate = pd.DataFrame(
        {
            "Date": date,
            "marketingExpenseRate": marketing_expense_rate,
            "administrativeExpenseRate": administrative_expense_rate,
            "RDExpenseRate": rd_expense_rate,
            "financialCostRate": financial_cost_rate,
            "depreciationCostRate": depreciation_cost_rate,
            "devaluationCostRate": devaluation_cost_rate,
            "incomeTaxRate": np.minimum(income_tax_rate, 0),
        }
    )

    earning_power = pd.DataFrame(
        {
            "Date": date,
            "grossProfitMargin": gross_profit_margin,
            "netProfitMargin": net_profit_margin,
            "ROE": roe,
            "ROA": roa,
            "ROIC": roic,
        }
    )

    safe_factor = pd.DataFrame(
        {
            "Date": date,
            "currentRatio": current_ratio,
            "quickRatio": quick_ratio,
        }
    )

    growth_rate = pd.DataFrame(
        {
            "Date": date,
            "operatingIncomeGrowthRate": operating_income_growth_rate,
            "netProfitGrowthRate": net_profit_growth_rate,
        }
    )

    operating_capability = pd.DataFrame(
        {
            "Date": date,
            "totalAssetsTurnOver": total_assets_turn_over,
            "fixedAssetsTurnOver": fixed_assets_turn_over,
            "inventoryTurnOver": inventory_turn_over,
            "accountsReceivableTurnOver": accounts_receivable_turn_over,
            "financeLeverageDegree": finance_leverage_degree,
        }
    )

    index_per_share = pd.DataFrame({"Date": date, "EPS": eps, "DEPS": deps, "DPS": dps})

    key_index = pd.DataFrame(
        {
            "代码": code,
            "名称": stock_name,
            "价格": round(last_price, 1),
            "总市值(亿)": round(market_value / 100000000, 1),
            "PE": round(pe_value, 1),
            "PB": round(adjusted_pb, 1),
            "ROE": round(safe_mean(roe, 3), 3),
            "ROIC": round(safe_mean(roic, 3), 3),
            "毛利率": round(safe_mean(gross_profit_margin, 3), 3),
            "净利率": round(safe_mean(net_profit_margin, 3), 3),
            "现金回报率": round(cash_return_rate, 3),
            "营收增长率": round(operating_income_cagr, 3),
            "净利润增长率": round(net_profit_cagr, 3),
        },
        index=[0],
    )

    report_tables = {
        "asset": asset_table,
        "balance": balance_table,
        "cash": cash_table,
        "operating_cash": operating_cash_table,
        "investment_cash": investment_cash_table,
        "financial_cash": financial_cash_table,
        "income": income_table,
        "earning_power": earning_power,
        "safe_factor": safe_factor,
        "growth_rate": growth_rate,
        "operating_capability": operating_capability,
        "expense_rate": expense_rate,
        "operating_funds": operating_funds,
        "index_per_share": index_per_share,
    }

    if write_artifacts:
        results_dir = project_root / "results"
        results_dir.mkdir(parents=True, exist_ok=True)

        save_report(
            results_dir / f"{code}_report.csv",
            [
                asset_table,
                balance_table,
                cash_table,
                operating_cash_table,
                investment_cash_table,
                financial_cash_table,
                income_table,
                earning_power,
                safe_factor,
                growth_rate,
                operating_capability,
            ],
        )

        if show_plot:
            plot_report(
                code=code,
                stock_name=stock_name,
                last_price=last_price,
                pb_value=adjusted_pb,
                pe_value=round(pe_value, 1),
                dv_value=dv_value,
                cash_return_rate=cash_return_rate,
                date=date,
                tables=report_tables,
                output_png=results_dir / f"{code}.png",
                show_plot=show_plot,
            )

    if return_details:
        return key_index, {
            "annual_metrics": _annual_metric_records(date, earning_power),
            "averages": key_index.iloc[0].to_dict(),
            "cash_return_rate": cash_return_rate,
        }
    return key_index


def write_results_to_sheet(summary_path: Path, sheet_name: str, results: pd.DataFrame) -> None:
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    if summary_path.exists():
        try:
            # Validate workbook first; corrupted files will be rotated.
            load_workbook(filename=summary_path)
        except (BadZipFile, KeyError, OSError, ValueError) as exc:
            backup = summary_path.with_suffix(
                summary_path.suffix + f".corrupt.{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            )
            summary_path.rename(backup)
            print(f"summary file is corrupted, moved to: {backup} ({exc})")
            with pd.ExcelWriter(summary_path, mode="w", engine="openpyxl") as writer:
                results.to_excel(writer, sheet_name=sheet_name, index=False)
            return
        try:
            with pd.ExcelWriter(
                summary_path,
                mode="a",
                engine="openpyxl",
                if_sheet_exists="replace",
            ) as writer:
                results.to_excel(writer, sheet_name=sheet_name, index=False)
        except TypeError:
            # Fallback for older pandas without if_sheet_exists.
            book = load_workbook(filename=summary_path)
            if sheet_name in book.sheetnames:
                del book[sheet_name]
            if not book.sheetnames:
                book.create_sheet("tmp")
            book.save(filename=summary_path)
            with pd.ExcelWriter(summary_path, mode="a", engine="openpyxl") as writer:
                results.to_excel(writer, sheet_name=sheet_name, index=False)
                if "tmp" in writer.book.sheetnames and len(writer.book.sheetnames) > 1:
                    del writer.book["tmp"]
    else:
        with pd.ExcelWriter(summary_path, mode="w", engine="openpyxl") as writer:
            results.to_excel(writer, sheet_name=sheet_name, index=False)


def main() -> None:
    args = parse_args()
    project_root = Path(__file__).resolve().parents[1]
    run_list = parse_code_inputs(args.codes)
    if not run_list:
        raise ValueError("No valid stock codes provided.")
    print(f"Run list: {run_list}")
    name_map = load_name_map_from_value_file(project_root, args.value_file, args.stats_sheet)
    print(f"Name map loaded from {args.value_file}#{args.stats_sheet}: {len(name_map)}")
    live_quotes = fetch_live_quote_metrics(
        run_list,
        timeout=args.quote_timeout,
        refresh=(not args.skip_live_quote_refresh),
    )

    skip_stocks = set()
    if args.download:
        for stock in run_list:
            print(f"正在下载财务报表: {stock}")
            try:
                getStockInfoWeb(stock)
            except Exception as exc:
                print(f"获取股票信息失败 {stock}: {exc}")
            try:
                for attempt in range(1, args.retry + 1):
                    downloadFinancialDataFutu(stock, yearCycle=True)
                    ok, reason = verify_code_integrity(project_root, stock, min_years=args.min_years)
                    if ok:
                        break
                    print(
                        f"数据完整性校验未通过 {stock} "
                        f"(attempt {attempt}/{args.retry}): {reason}"
                    )
                ok, reason = verify_code_integrity(project_root, stock, min_years=args.min_years)
                if not ok:
                    print(f"下载后仍不完整 {stock}: {reason}")
                    skip_stocks.add(stock)
            except Exception as exc:
                print(f"下载财务报表失败 {stock}: {exc}")
                # Fallback to existing local data if integrity is still acceptable.
                try:
                    ok, reason = verify_code_integrity(project_root, stock, min_years=args.min_years)
                    if not ok:
                        print(f"本地历史数据也不完整，跳过分析 {stock}: {reason}")
                        skip_stocks.add(stock)
                except Exception as exc2:
                    print(f"无法使用本地历史数据，跳过分析 {stock}: {exc2}")
                    skip_stocks.add(stock)

    print("开始分析财务报表...")
    results_frames: List[pd.DataFrame] = []
    for stock in run_list:
        if stock in skip_stocks:
            print(f"跳过分析 {stock}: 数据不完整")
            continue
        try:
            key_index = financial_analysis(
                project_root,
                stock,
                show_plot=args.plot,
                max_years=args.max_years,
                name_map=name_map,
                live_quotes=live_quotes,
                write_artifacts=True,
            )
            results_frames.append(key_index)
        except Exception as exc:
            print(f"分析失败 {stock}: {exc}")

    if not results_frames:
        raise RuntimeError("No stocks analyzed successfully.")

    results = pd.concat(results_frames, ignore_index=True)
    print(results)
    summary_path = project_root / args.summary_file
    write_results_to_sheet(summary_path, args.sheet, results)
    print(f"Summary saved to: {summary_path} (sheet={args.sheet})")


if __name__ == "__main__":
    main()
