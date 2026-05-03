#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class MetricSeries:
    label: str
    values: np.ndarray
    rows_used: List[str]
    kind: str = "amount"  # amount | price | ratio


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Market value evaluation sheet generator")
    parser.add_argument("--code", default="HK.00700", help="Stock code for downloaded financial data")
    parser.add_argument("--company", default="", help="Company name shown in output sheet (auto from info if empty)")
    parser.add_argument(
        "--workbook",
        default=str(Path(__file__).resolve().parents[1] / "data" / "Value_Investment.xlsx"),
        help="Target Excel workbook path",
    )
    parser.add_argument("--sheet", default="", help="Output sheet name (default: <company>_新估值)")
    parser.add_argument(
        "--years",
        default="auto",
        help="Comma-separated fiscal years, or 'auto' for latest 5 common years",
    )
    parser.add_argument("--overwrite", action="store_true", help="Replace sheet if it already exists")
    parser.add_argument(
        "--fx-rate",
        type=float,
        default=None,
        help="Manual FX rate: 1 unit of report currency equals N units of quote currency",
    )
    parser.add_argument("--report-currency", default="", help="Force report currency (CNY/HKD/USD/JPY etc.)")
    parser.add_argument("--quote-currency", default="", help="Force quote currency (CNY/HKD/USD/JPY etc.)")
    parser.add_argument(
        "--yahoo-timeout",
        type=int,
        default=10,
        help="Yahoo HTTP timeout (seconds) for shares-outstanding fallback",
    )
    return parser.parse_args()


def parse_number(value: object) -> float:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", "").replace(" ", "")
    if text in {"", "--", "---", "nan", "None"}:
        return 0.0

    if "%" in text:
        tail = text.rsplit("%", 1)[-1]
        if re.search(r"\d", tail):
            text = tail

    text = re.sub(r"^--+(?=\d)", "-", text)

    m = re.findall(r"-?\d+(?:\.\d+)?", text)
    if not m:
        return 0.0

    number = float(m[0])
    if "萬億" in text:
        number *= 1e12
    elif "億" in text:
        number *= 1e8
    elif "萬" in text:
        number *= 1e4
    return number


def _normalize_code(raw_code: str) -> str:
    text = str(raw_code).strip().upper()
    if not text:
        return text
    if "." in text:
        return text
    if re.fullmatch(r"\d{6}", text):
        if text.startswith(("6", "9")):
            return f"SH.{text}"
        if text.startswith(("0", "3")):
            return f"SZ.{text}"
    if re.fullmatch(r"\d{1,5}", text):
        return f"HK.{text.zfill(5)}"
    return f"US.{text}"


def _normalize_ccy(text: str) -> str:
    raw = text.strip().upper()
    if raw in {"CNY", "RMB", "CNH", "HKD", "USD", "JPY", "EUR", "GBP"}:
        return "CNY" if raw in {"RMB", "CNH"} else raw

    mapping = {
        "人民币": "CNY",
        "人民幣": "CNY",
        "港币": "HKD",
        "港幣": "HKD",
        "美元": "USD",
        "美金": "USD",
        "日元": "JPY",
        "日圆": "JPY",
        "日圓": "JPY",
        "欧元": "EUR",
        "歐元": "EUR",
        "英镑": "GBP",
        "英鎊": "GBP",
    }
    return mapping.get(text.strip(), raw)


def _display_ccy(ccy: str) -> str:
    return {
        "CNY": "人民币",
        "HKD": "港币",
        "USD": "美元",
        "JPY": "日元",
        "EUR": "欧元",
        "GBP": "英镑",
    }.get(ccy, ccy)


def _quote_currency_from_code(code: str) -> str:
    prefix = code.split(".", 1)[0].upper() if "." in code else ""
    return {
        "SH": "CNY",
        "SZ": "CNY",
        "HK": "HKD",
        "US": "USD",
        "JP": "JPY",
    }.get(prefix, "CNY")


def _yahoo_symbol_from_code(code: str) -> Optional[str]:
    c = code.strip().upper()
    if "." not in c:
        return None
    market, ticker = c.split(".", 1)
    if market == "HK":
        if ticker.isdigit():
            return f"{int(ticker):04d}.HK"
        return f"{ticker}.HK"
    if market == "SH":
        return f"{ticker}.SS"
    if market == "SZ":
        return f"{ticker}.SZ"
    if market == "US":
        return ticker
    if market == "JP":
        return f"{ticker}.T"
    return None


def _first_currency_from_frame(df: pd.DataFrame) -> Optional[str]:
    index_texts = [str(x).strip() for x in df.index.tolist()]

    # 0) Direct currency tokens in index rows (some tables put currency as its own row label).
    for idx in index_texts:
        ccy = _normalize_ccy(idx)
        if ccy in {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}:
            return ccy

    marker_rows = ["貨幣單位", "货币单位", "幣種", "币种", "Currency"]
    for row in marker_rows:
        if row in df.index:
            # 1) Currency value stored in same row cells.
            for val in df.loc[row].tolist():
                ccy = _normalize_ccy(str(val))
                if ccy in {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}:
                    return ccy

            # 2) Currency stored in nearby index rows (e.g., marker row followed by "美元").
            pos = index_texts.index(row)
            for j in range(max(0, pos - 2), min(len(index_texts), pos + 5)):
                ccy = _normalize_ccy(index_texts[j])
                if ccy in {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}:
                    return ccy

    # 3) Generic marker contains check + nearby scan.
    for i, idx in enumerate(index_texts):
        low = idx.lower()
        if any(k in low for k in ["貨幣", "货币", "幣種", "币种", "currency"]):
            for j in range(max(0, i - 2), min(len(index_texts), i + 5)):
                ccy = _normalize_ccy(index_texts[j])
                if ccy in {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}:
                    return ccy
    return None


def _detect_report_currency(income_df: pd.DataFrame, balance_df: pd.DataFrame, cash_df: pd.DataFrame) -> str:
    for df in [income_df, balance_df, cash_df]:
        c = _first_currency_from_frame(df)
        if c:
            return c
    return "CNY"


def _report_currency_marked(income_df: pd.DataFrame, balance_df: pd.DataFrame, cash_df: pd.DataFrame) -> bool:
    return any(_first_currency_from_frame(df) is not None for df in [income_df, balance_df, cash_df])


def _sheet_saved_report_currency(workbook_path: Path, sheet_name: str) -> Optional[str]:
    if not workbook_path.exists():
        return None
    try:
        wb = load_workbook(workbook_path, data_only=True)
    except Exception:
        return None
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    if str(ws["A3"].value or "").strip() != "Report Currency":
        return None
    ccy = _normalize_ccy(str(ws["B3"].value or "").strip())
    if ccy in {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}:
        return ccy
    return None


def _year_from_col(col: object) -> Optional[int]:
    m = re.match(r"^(\d{4})/\d{2}/\d{2}(?:\.\d+)?$", str(col))
    if not m:
        return None
    return int(m.group(1))


def _choose_columns_by_year(df: pd.DataFrame) -> pd.DataFrame:
    chosen: Dict[int, Tuple[int, str]] = {}
    for col in df.columns:
        col_text = str(col)
        year = _year_from_col(col_text)
        if year is None:
            continue
        # Prefer unsuffixed column like 2024/12/31 over 2024/12/31.1
        score = 0 if re.match(r"^\d{4}/\d{2}/\d{2}$", col_text) else 1
        if year not in chosen or score < chosen[year][0]:
            chosen[year] = (score, col_text)

    years = sorted(chosen.keys(), reverse=True)
    cols = [chosen[y][1] for y in years]
    out = df[cols].copy()
    out.columns = years
    return out


def _infer_common_years(
    income_df: pd.DataFrame,
    balance_df: pd.DataFrame,
    cash_df: pd.DataFrame,
    count: int = 5,
) -> List[int]:
    iy = set(_choose_columns_by_year(income_df).columns.tolist())
    by = set(_choose_columns_by_year(balance_df).columns.tolist())
    cy = set(_choose_columns_by_year(cash_df).columns.tolist())
    common = sorted(iy & by & cy, reverse=True)
    if not common:
        raise ValueError("No common fiscal years across income/balance/cash")
    picked = common[:count]
    return sorted(picked)


def values_for_row(df_by_year: pd.DataFrame, row_name: str, years: Sequence[int]) -> np.ndarray:
    if row_name not in df_by_year.index:
        return np.zeros(len(years), dtype=float)
    vals = []
    for y in years:
        if y in df_by_year.columns:
            vals.append(parse_number(df_by_year.loc[row_name, y]))
        else:
            vals.append(0.0)
    return np.array(vals, dtype=float)


def first_existing_series(
    df_by_year: pd.DataFrame,
    row_candidates: Sequence[str],
    years: Sequence[int],
    label: str,
) -> MetricSeries:
    for row in row_candidates:
        if row in df_by_year.index:
            return MetricSeries(label=label, values=values_for_row(df_by_year, row, years), rows_used=[row])
    return MetricSeries(label=label, values=np.zeros(len(years), dtype=float), rows_used=[])


def sum_unique_rows(
    df_by_year: pd.DataFrame,
    row_candidates: Sequence[str],
    years: Sequence[int],
    used_rows: set,
    label: str,
) -> MetricSeries:
    total = np.zeros(len(years), dtype=float)
    used: List[str] = []
    for row in row_candidates:
        if row in df_by_year.index and row not in used_rows:
            total += values_for_row(df_by_year, row, years)
            used.append(row)
            used_rows.add(row)
    return MetricSeries(label=label, values=total, rows_used=used)


def _fx_from_stats_sheet(workbook_path: Path, from_ccy: str, to_ccy: str) -> Optional[Tuple[float, str]]:
    if from_ccy == to_ccy:
        return 1.0, "same currency"
    if not workbook_path.exists():
        return None

    wb = load_workbook(workbook_path, data_only=True)
    if "统计" not in wb.sheetnames:
        return None
    ws = wb["统计"]

    valid_ccy = {"CNY", "HKD", "USD", "JPY", "EUR", "GBP"}
    rates: Dict[Tuple[str, str], float] = {}

    # Detect any row that looks like a currency header row (contains >=2 currency codes).
    header_rows: Dict[int, Dict[int, str]] = {}
    for r in range(1, ws.max_row + 1):
        row_map: Dict[int, str] = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                ccy = _normalize_ccy(v)
                if ccy in valid_ccy:
                    row_map[c] = ccy
        if len(row_map) >= 2:
            header_rows[r] = row_map

    for header_row, col_ccy in header_rows.items():
        min_header_col = min(col_ccy.keys())
        label_cols = [x for x in [min_header_col - 1, min_header_col - 2] if x >= 1]
        for r in range(1, ws.max_row + 1):
            if r == header_row:
                continue

            row_ccy = ""
            for lc in label_cols:
                v = ws.cell(r, lc).value
                if isinstance(v, str):
                    ccy = _normalize_ccy(v)
                    if ccy in valid_ccy:
                        row_ccy = ccy
                        break
            if not row_ccy:
                continue

            has_numeric = False
            for c, ccy_to in col_ccy.items():
                num = parse_number(ws.cell(r, c).value)
                if num > 0:
                    has_numeric = True
                    rates[(row_ccy, ccy_to)] = num
            if not has_numeric:
                continue

    if (from_ccy, to_ccy) in rates:
        return rates[(from_ccy, to_ccy)], "stats-sheet direct"
    if (to_ccy, from_ccy) in rates and rates[(to_ccy, from_ccy)] > 0:
        return 1.0 / rates[(to_ccy, from_ccy)], "stats-sheet inverted"

    # Multi-hop inference (e.g., JPY->HKD->USD).
    graph: Dict[str, Dict[str, float]] = {}
    for (a, b), r in rates.items():
        if r <= 0:
            continue
        graph.setdefault(a, {})[b] = r
        graph.setdefault(b, {})[a] = 1.0 / r

    if from_ccy in graph and to_ccy in graph:
        queue: List[Tuple[str, float, List[str]]] = [(from_ccy, 1.0, [from_ccy])]
        visited = {from_ccy}
        while queue:
            cur, rate_so_far, path = queue.pop(0)
            if len(path) > 5:
                continue
            for nxt, edge_rate in graph.get(cur, {}).items():
                if nxt in path:
                    continue
                new_rate = rate_so_far * edge_rate
                new_path = path + [nxt]
                if nxt == to_ccy:
                    return new_rate, f"stats-sheet path {'->'.join(new_path)}"
                if nxt not in visited:
                    visited.add(nxt)
                    queue.append((nxt, new_rate, new_path))

    return None


def _resolve_fx_rate(
    workbook_path: Path,
    report_currency: str,
    quote_currency: str,
    manual_fx: Optional[float],
) -> Tuple[float, str]:
    if quote_currency == report_currency:
        return 1.0, "same currency"
    if manual_fx and manual_fx > 0:
        return float(manual_fx), "manual --fx-rate"

    from_sheet = _fx_from_stats_sheet(workbook_path, report_currency, quote_currency)
    if from_sheet:
        return from_sheet

    raise ValueError(
        f"No FX rate found for {report_currency}->{quote_currency}. "
        "Please pass --fx-rate (1 report currency = N quote currency)."
    )


def _shares_from_stats_sheet(code: str, workbook_path: Optional[Path]) -> float:
    if workbook_path is None or (not workbook_path.exists()):
        return 0.0
    try:
        wb = load_workbook(workbook_path, data_only=True)
    except Exception:
        return 0.0
    if "统计" not in wb.sheetnames:
        return 0.0
    ws = wb["统计"]

    code_col = None
    shares_col = None
    for c in range(1, ws.max_column + 1):
        head = str(ws.cell(1, c).value or "").strip()
        if head == "代码":
            code_col = c
        elif head == "总股本":
            shares_col = c
    if code_col is None or shares_col is None:
        return 0.0

    target = _normalize_code(code)
    for r in range(2, ws.max_row + 1):
        cell_code = _normalize_code(str(ws.cell(r, code_col).value or "").strip())
        if not cell_code or cell_code != target:
            continue
        v = parse_number(ws.cell(r, shares_col).value)
        if v <= 0:
            continue
        # In this sheet, total shares is usually stored in 亿股.
        if v < 1e7:
            return v * 1e8
        return v
    return 0.0


def _info_value(info_df: pd.DataFrame, keys: Sequence[str]) -> float:
    for key in keys:
        if key in info_df.index:
            row = info_df.loc[key]
            if isinstance(row, pd.Series):
                vals = row.tolist()
            else:
                vals = [row]
            for v in vals:
                num = parse_number(v)
                if num > 0:
                    return num
    return 0.0


def _fetch_shares_from_yahoo(code: str, timeout: int = 10) -> float:
    symbol = _yahoo_symbol_from_code(code)
    if not symbol:
        return 0.0
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
        )
    }

    def _raw(v: object) -> float:
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, dict) and isinstance(v.get("raw"), (int, float)):
            return float(v.get("raw"))
        return 0.0

    urls = [
        f"https://query1.finance.yahoo.com/v10/finance/quoteSummary/{symbol}?modules=defaultKeyStatistics,price,summaryDetail",
        f"https://query1.finance.yahoo.com/v7/finance/quote?symbols={symbol}",
    ]
    for url in urls:
        try:
            resp = requests.get(url, headers=headers, timeout=timeout)
            if resp.status_code != 200:
                continue
            data = resp.json()
            # quoteSummary payload
            result = (((data.get("quoteSummary") or {}).get("result")) or [None])[0]
            if isinstance(result, dict):
                dks = result.get("defaultKeyStatistics") or {}
                price = result.get("price") or {}
                sd = result.get("summaryDetail") or {}
                shares = _raw(dks.get("sharesOutstanding")) or _raw(price.get("sharesOutstanding")) or _raw(
                    sd.get("sharesOutstanding")
                )
                if shares > 0:
                    return shares
                market_cap = _raw(price.get("marketCap")) or _raw(sd.get("marketCap"))
                px = _raw(price.get("regularMarketPrice"))
                if market_cap > 0 and px > 0:
                    return market_cap / px
            # quote payload
            result = ((((data.get("quoteResponse") or {}).get("result")) or [None])[0])
            if isinstance(result, dict):
                shares = _raw(result.get("sharesOutstanding"))
                if shares > 0:
                    return shares
                market_cap = _raw(result.get("marketCap"))
                px = _raw(result.get("regularMarketPrice"))
                if market_cap > 0 and px > 0:
                    return market_cap / px
        except Exception:
            continue
    return 0.0


def _resolve_total_shares(
    info_df: pd.DataFrame,
    code: str,
    yahoo_timeout: int = 10,
    shares_fallback_workbook: Optional[Path] = None,
) -> Tuple[float, str]:
    market_value = _info_value(info_df, ["总市值", "總市值"])
    last_price = _info_value(info_df, ["平均价", "最新价", "現價", "现价"])
    implied_shares = (market_value / last_price) if (market_value > 0 and last_price > 0) else 0.0

    total_shares = _info_value(info_df, ["总股本", "總股本", "Shares Outstanding"])
    if total_shares > 0:
        # Some sources may return plain "亿股" without unit; normalize when obviously too small.
        if total_shares < 1e7:
            scaled = total_shares * 1e8
            if scaled > 0:
                total_shares = scaled
        if implied_shares > 0:
            ratio = total_shares / implied_shares if implied_shares > 0 else 1.0
            if ratio < 0.5 or ratio > 2.0:
                return implied_shares, "Futu market cap / price (override inconsistent info shares)"
        return total_shares, "Futu total shares"

    from_stats = _shares_from_stats_sheet(code, shares_fallback_workbook)
    if from_stats > 0:
        if implied_shares > 0:
            ratio = from_stats / implied_shares if implied_shares > 0 else 1.0
            if ratio < 0.5 or ratio > 2.0:
                return implied_shares, "Futu market cap / price (override inconsistent stats shares)"
        return from_stats, "Stats sheet total shares"

    yahoo_shares = _fetch_shares_from_yahoo(code, timeout=yahoo_timeout)
    if yahoo_shares > 0:
        return yahoo_shares, "Yahoo sharesOutstanding"

    if implied_shares > 0:
        return implied_shares, "Futu market cap / price"
    raise ValueError(f"Unable to resolve total shares for {code}")


def build_sheet_rows(
    income_df_raw: pd.DataFrame,
    balance_df_raw: pd.DataFrame,
    cash_df_raw: pd.DataFrame,
    info_df: pd.DataFrame,
    code: str,
    years: Sequence[int],
    fx_rate: float,
    yahoo_timeout: int,
    shares_fallback_workbook: Optional[Path] = None,
) -> Tuple[List[MetricSeries], Dict[str, float]]:
    income_df = _choose_columns_by_year(income_df_raw)
    balance_df = _choose_columns_by_year(balance_df_raw)
    cash_df = _choose_columns_by_year(cash_df_raw)

    total_revenue = first_existing_series(
        income_df,
        ["營業總收入", "营业总收入", "總收入", "总收入", "營業收入", "营业收入", "Revenue", "Total Revenue"],
        years,
        "Revenue",
    )

    operating_profit = first_existing_series(
        income_df,
        ["經營溢利", "營運溢利", "營業利潤", "營業利益", "Operating Profit", "EBIT"],
        years,
        "Operating Profit",
    )

    selling_expense = first_existing_series(
        income_df,
        ["銷售費用", "销售费用", "Selling Expense", "Selling Expenses"],
        years,
        "Selling Expenses",
    )
    admin_expense = first_existing_series(
        income_df,
        ["行政費用", "管理费用", "Administrative Expense", "General and Administrative Expense"],
        years,
        "Administrative Expenses",
    )
    rd_expense = first_existing_series(
        income_df,
        ["研發費用", "研发费用", "研發開支", "Research and Development", "R&D Expense"],
        years,
        "R&D Expenses",
    )

    if not selling_expense.rows_used and not admin_expense.rows_used:
        sga = first_existing_series(
            income_df,
            ["銷售和管理費用", "銷售及管理費用", "SG&A", "Selling, General and Administrative Expense"],
            years,
            "Selling Expenses",
        )
        selling_expense = sga

    pre_tax_profit = first_existing_series(
        income_df,
        ["除稅前溢利", "税前利润", "稅前利潤", "Pre-Tax Profit", "Profit Before Tax"],
        years,
        "Pre-Tax Profit",
    )
    tax_expense = first_existing_series(
        income_df,
        ["稅項", "所得税费用", "税项", "所得稅", "Income Tax"],
        years,
        "Tax",
    )

    dep_amort = first_existing_series(
        cash_df,
        ["折舊及攤銷:", "折舊和攤銷", "Depreciation and Amortization"],
        years,
        "Depreciation & Amortization",
    )
    if not dep_amort.rows_used:
        dep = first_existing_series(cash_df, ["-折舊", "折舊"], years, "Depreciation")
        amo = first_existing_series(cash_df, ["-無形資產攤銷", "攤銷"], years, "Amortization")
        other = first_existing_series(cash_df, ["-其他折舊及攤銷"], years, "Other D&A")
        dep_amort = MetricSeries(
            label="Depreciation & Amortization",
            values=dep.values + amo.values + other.values,
            rows_used=dep.rows_used + amo.rows_used + other.rows_used,
        )

    tax_rate_series = np.divide(
        np.abs(tax_expense.values),
        np.abs(pre_tax_profit.values),
        out=np.zeros(len(years), dtype=float),
        where=np.abs(pre_tax_profit.values) > 1e-12,
    )
    tax_rate_nonzero = tax_rate_series[tax_rate_series > 0]
    avg_tax_rate = float(np.mean(tax_rate_nonzero)) if tax_rate_nonzero.size else 0.2

    adjusted_operating_profit = operating_profit.values + 0.1 * (
        np.abs(selling_expense.values) + np.abs(admin_expense.values) + np.abs(rd_expense.values)
    )
    estimated_net_profit = adjusted_operating_profit * (1 - avg_tax_rate)
    adjusted_net_profit = estimated_net_profit + 0.1 * np.abs(dep_amort.values)

    reported_net_profit = first_existing_series(
        income_df,
        [
            "股東應占溢利",
            "股东应占溢利",
            "本公司股东应占溢利",
            "本公司股東應佔溢利",
            "归属于母公司股东的净利润",
            "归属于母公司股东应占溢利",
            "歸屬於母公司股東應佔溢利",
            "母公司擁有人應佔溢利",
            "净利润",
            "淨利潤",
            "年内利润",
            "Net Profit",
            "Net Income",
        ],
        years,
        "Net Profit",
    )

    lower_mv_core = adjusted_net_profit * 12.5
    upper_mv_core = adjusted_net_profit * 33.3

    add_items = [
        (
            "Cash and Equivalents",
            ["現金及等價物", "現金及現金等價物", "-現金和現金等價物", "現金及现金等價物", "现金及现金等价物", "貨幣資金", "货币资金"],
        ),
        ("Short-Term Deposit", ["短期存款", "短期存单", "短期存單"]),
        (
            "Current and Non-Current Fixed Deposit",
            [
                "定期存款-流動資產",
                "定期存款-非流動資產",
                "定期存款",
                "一年內到期的定期存款",
                "非流動定期存款",
            ],
        ),
        (
            "Short-Term and Long-Term Investments",
            [
                "短期投資",
                "短期投资",
                "證券投資",
                "证券投资",
                "其他投資",
                "其他投资",
                "-短期投資",
                "-短期投资",
            ],
        ),
        (
            "Financial Assets at Fair Value",
            [
                "按公平值入損益金融資產-流動資產",
                "按公平值入損益金融資產-非流動資產",
                "按公允價值計量且其變動計入當期損益的金融資產",
                "按公允價值計量且其變動計入當期損益的金融資產-流動資產",
                "按公允價值計量且其變動計入當期損益的金融資產-非流動資產",
                "按公允价值计量且其变动计入当期损益的金融资产",
                "按公允价值计量且其变动计入当期损益的金融资产-流动资产",
                "按公允价值计量且其变动计入当期损益的金融资产-非流动资产",
                "交易性金融資產",
                "交易性金融资产",
                "其他非流動金融資產",
                "其他非流动金融资产",
                "其他權益工具投資",
                "其他权益工具投资",
                "可供出售金融資產-流動資產",
                "可供出售金融資產-非流動資產",
                "可供出售金融资产-流动资产",
                "可供出售金融资产-非流动资产",
            ],
        ),
        ("Secured Deposit", ["已抵押存款", "已抵押存款-非流動資產", "受限制貨幣資金", "受限制货币资金"]),
        ("Investment Property", ["投資物業", "投资物业", "投資性房地產", "投资性房地产"]),
        ("Long-Term Investment", ["長期投資", "长期投资", "長期股權投資", "长期股权投资"]),
        ("Associated Company Interest", ["聯營公司權益", "關聯公司權益", "联营公司权益", "對聯營企業投資", "对联营企业投资"]),
        ("Joint Control of Entity Interests", ["共同控制實體權益", "共同控制实体权益"]),
        ("Interests in Joint Venture", ["合營公司權益", "合营公司权益", "於合營企業投資", "于合营企业投资", "於聯營和合營公司投資", "于联营和合营公司投资"]),
    ]

    subtract_items = [
        ("Short-Term Debt", ["短期借款", "-短期借款", "短期借款與租賃負債"]),
        ("Bank Loans and Overdrafts", ["銀行貸款及透支", "Bank Loans and Overdrafts"]),
        ("Long-Term Bank Loan", ["長期銀行貸款", "长期银行贷款", "-長期借款", "-长期借款", "長期借款與租賃負債", "长期借款与租赁负债"]),
    ]

    used_asset_rows: set = set()
    add_series: List[MetricSeries] = []
    for label, rows in add_items:
        add_series.append(sum_unique_rows(balance_df, rows, years, used_asset_rows, label))

    add_total = np.sum([x.values for x in add_series], axis=0) if add_series else np.zeros(len(years), dtype=float)

    # If "現金及現金等價物和短期投資" exists and sub-items are missing, use it once.
    combined_cash_sti = first_existing_series(
        balance_df,
        ["現金及現金等價物和短期投資", "现金及现金等价物和短期投资", "貨幣資金及交易性金融資產", "货币资金及交易性金融资产"],
        years,
        "Combined Cash+Short Investment",
    )
    if combined_cash_sti.rows_used and np.allclose(add_total, 0):
        add_total += combined_cash_sti.values
        add_series[0] = MetricSeries(
            label=add_series[0].label,
            values=combined_cash_sti.values,
            rows_used=combined_cash_sti.rows_used,
        )

    used_debt_rows: set = set()
    subtract_series: List[MetricSeries] = []
    for label, rows in subtract_items:
        subtract_series.append(sum_unique_rows(balance_df, rows, years, used_debt_rows, label))

    subtract_total = (
        np.sum([x.values for x in subtract_series], axis=0) if subtract_series else np.zeros(len(years), dtype=float)
    )

    lower_mv_total_report = lower_mv_core + add_total - subtract_total
    upper_mv_total_report = upper_mv_core + add_total - subtract_total

    total_shares, shares_source = _resolve_total_shares(
        info_df,
        code=code,
        yahoo_timeout=yahoo_timeout,
        shares_fallback_workbook=shares_fallback_workbook,
    )

    lower_price = np.divide(
        lower_mv_total_report,
        total_shares,
        out=np.zeros(len(years), dtype=float),
        where=total_shares > 1e-12,
    )
    upper_price = np.divide(
        upper_mv_total_report,
        total_shares,
        out=np.zeros(len(years), dtype=float),
        where=total_shares > 1e-12,
    )

    # Keep all calculations in report currency (raw financial statement currency).
    rows: List[MetricSeries] = [
        MetricSeries(total_revenue.label, total_revenue.values, total_revenue.rows_used),
        MetricSeries("Revenue YoY Growth", np.zeros(len(years), dtype=float), [], kind="ratio"),
        MetricSeries(operating_profit.label, operating_profit.values, operating_profit.rows_used),
        MetricSeries(selling_expense.label, selling_expense.values, selling_expense.rows_used),
        MetricSeries(admin_expense.label, admin_expense.values, admin_expense.rows_used),
        MetricSeries(rd_expense.label, rd_expense.values, rd_expense.rows_used),
        MetricSeries("Adjusted Operating Profit (+10% S&A+R&D)", adjusted_operating_profit, []),
        MetricSeries("Estimated Net Profit (avg tax)", estimated_net_profit, []),
        MetricSeries(dep_amort.label, dep_amort.values, dep_amort.rows_used),
        MetricSeries(reported_net_profit.label, reported_net_profit.values, reported_net_profit.rows_used),
        MetricSeries("Adjusted Net Profit (+10% D&A)", adjusted_net_profit, []),
        MetricSeries("Lowest Market Value Core (x12.5)", lower_mv_core, []),
        MetricSeries("Highest Market Value Core (x33.3)", upper_mv_core, []),
    ]

    for s in add_series:
        rows.append(MetricSeries(s.label, s.values, s.rows_used))
    rows.append(MetricSeries("Total Additions", add_total, []))

    for s in subtract_series:
        rows.append(MetricSeries(s.label, s.values, s.rows_used))
    rows.append(MetricSeries("Total Subtractions", subtract_total, []))

    rows.extend(
        [
            MetricSeries("Lowest Market Value (after add/sub)", lower_mv_total_report, []),
            MetricSeries("Highest Market Value (after add/sub)", upper_mv_total_report, []),
            MetricSeries("Lower Price Limit", lower_price, [], kind="price"),
            MetricSeries("Upper Price Limit", upper_price, [], kind="price"),
        ]
    )

    summary = {
        "avg_tax_rate": avg_tax_rate,
        "total_shares": total_shares,
        "shares_source": shares_source,
    }
    return rows, summary


def write_sheet(
    workbook_path: Path,
    sheet_name: str,
    company: str,
    years: Sequence[int],
    rows: Sequence[MetricSeries],
    summary: Dict[str, float],
    overwrite: bool,
    report_currency: str,
    quote_currency: str,
    fx_rate: float,
    fx_source: str,
) -> None:
    wb = load_workbook(workbook_path)
    if sheet_name in wb.sheetnames:
        if not overwrite:
            raise ValueError(f"Sheet '{sheet_name}' already exists. Use --overwrite to replace.")
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name)
    ws["A1"] = f"{company} Market Value Evaluation (Auto)"
    ws["A2"] = "Method: Adjust OP -> Tax -> Adjusted NP -> x12.5/x33.3 -> Add assets -> Subtract debts -> /Shares"
    ws["A3"] = "Report Currency"
    ws["B3"] = report_currency
    ws["A4"] = "Quote Currency"
    ws["B4"] = quote_currency
    ws["A5"] = "FX (report->quote)"
    ws["B5"] = fx_rate
    ws["C5"] = fx_source
    ws["A6"] = "Average Tax Rate"
    ws["B6"] = summary["avg_tax_rate"]
    ws["B6"].number_format = "0.00%"
    ws["A7"] = "Total Shares"
    ws["B7"] = summary["total_shares"] / 1e8
    ws["B7"].number_format = "0.00"
    ws["C7"] = "亿股"

    # Tunable parameters (editable in Excel).
    ws["J3"] = "S&A&R&D Add-back Ratio"
    ws["K3"] = 0.10
    ws["K3"].number_format = "0.00%"
    ws["J4"] = "D&A Add-back Ratio"
    ws["K4"] = 0.10
    ws["K4"].number_format = "0.00%"
    ws["J5"] = "Low EPV Multiple"
    ws["K5"] = 12.5
    ws["K5"].number_format = "0.00"
    ws["J6"] = "High EPV Multiple"
    ws["K6"] = 33.3
    ws["K6"].number_format = "0.00"

    start_row = 9
    ws.cell(start_row, 1, "Item")
    for i, y in enumerate(years, start=2):
        ws.cell(start_row, i, y)
    ws.cell(start_row, len(years) + 2, "Unit")
    ws.cell(start_row, len(years) + 3, "Source Rows")

    amount_unit = f"亿{_display_ccy(report_currency)}"
    price_unit = f"{report_currency}/股"

    r = start_row + 1
    label_rows: Dict[str, int] = {}
    for metric in rows:
        ws.cell(r, 1, metric.label)
        label_rows[metric.label] = r

        values = metric.values.copy()
        if metric.kind == "price":
            unit = price_unit
        elif metric.kind == "ratio":
            unit = "%"
        else:
            values = values / 1e8
            unit = amount_unit

        for i, v in enumerate(values, start=2):
            cell = ws.cell(r, i)
            if metric.kind == "ratio":
                cell.value = round(float(v), 4)
                cell.number_format = "0.00%"
            else:
                cell.value = round(float(v), 2)
                cell.number_format = "0.00"
        ws.cell(r, len(years) + 2, unit)
        ws.cell(r, len(years) + 3, ", ".join(metric.rows_used))
        r += 1

    # Put core valuation logic into Excel formulas for easy adjustment.
    required_labels = [
        "Revenue",
        "Revenue YoY Growth",
        "Operating Profit",
        "Selling Expenses",
        "Administrative Expenses",
        "R&D Expenses",
        "Adjusted Operating Profit (+10% S&A+R&D)",
        "Estimated Net Profit (avg tax)",
        "Depreciation & Amortization",
        "Adjusted Net Profit (+10% D&A)",
        "Lowest Market Value Core (x12.5)",
        "Highest Market Value Core (x33.3)",
        "Cash and Equivalents",
        "Interests in Joint Venture",
        "Total Additions",
        "Short-Term Debt",
        "Long-Term Bank Loan",
        "Total Subtractions",
        "Lowest Market Value (after add/sub)",
        "Highest Market Value (after add/sub)",
        "Lower Price Limit",
        "Upper Price Limit",
    ]
    missing_labels = [x for x in required_labels if x not in label_rows]
    if missing_labels:
        raise ValueError(f"Cannot locate formula rows in sheet: {missing_labels}")

    rev_row = label_rows["Revenue"]
    rev_growth_row = label_rows["Revenue YoY Growth"]
    op_row = label_rows["Operating Profit"]
    sell_row = label_rows["Selling Expenses"]
    admin_row = label_rows["Administrative Expenses"]
    rd_row = label_rows["R&D Expenses"]
    adj_op_row = label_rows["Adjusted Operating Profit (+10% S&A+R&D)"]
    est_np_row = label_rows["Estimated Net Profit (avg tax)"]
    da_row = label_rows["Depreciation & Amortization"]
    adj_np_row = label_rows["Adjusted Net Profit (+10% D&A)"]
    low_core_row = label_rows["Lowest Market Value Core (x12.5)"]
    high_core_row = label_rows["Highest Market Value Core (x33.3)"]
    add_start_row = label_rows["Cash and Equivalents"]
    add_end_row = label_rows["Interests in Joint Venture"]
    add_total_row = label_rows["Total Additions"]
    sub_start_row = label_rows["Short-Term Debt"]
    sub_end_row = label_rows["Long-Term Bank Loan"]
    sub_total_row = label_rows["Total Subtractions"]
    low_mv_row = label_rows["Lowest Market Value (after add/sub)"]
    high_mv_row = label_rows["Highest Market Value (after add/sub)"]
    low_px_row = label_rows["Lower Price Limit"]
    high_px_row = label_rows["Upper Price Limit"]

    for i in range(len(years)):
        col = get_column_letter(2 + i)
        if i == 0:
            ws[f"{col}{rev_growth_row}"] = ""
        else:
            prev_col = get_column_letter(1 + i)
            ws[f"{col}{rev_growth_row}"] = (
                f'=IF(ABS({prev_col}{rev_row})<1E-12,"",ROUND({col}{rev_row}/{prev_col}{rev_row}-1,4))'
            )
        ws[f"{col}{rev_growth_row}"].number_format = "0.00%"
        ws[f"{col}{adj_op_row}"] = (
            f"=ROUND({col}{op_row}+ABS({col}{sell_row})*$K$3+ABS({col}{admin_row})*$K$3+ABS({col}{rd_row})*$K$3,2)"
        )
        ws[f"{col}{est_np_row}"] = f"=ROUND({col}{adj_op_row}*(1-$B$6),2)"
        ws[f"{col}{adj_np_row}"] = f"=ROUND({col}{est_np_row}+ABS({col}{da_row})*$K$4,2)"
        ws[f"{col}{low_core_row}"] = f"=ROUND({col}{adj_np_row}*$K$5,2)"
        ws[f"{col}{high_core_row}"] = f"=ROUND({col}{adj_np_row}*$K$6,2)"
        ws[f"{col}{add_total_row}"] = f"=ROUND(SUM({col}{add_start_row}:{col}{add_end_row}),2)"
        ws[f"{col}{sub_total_row}"] = f"=ROUND(SUM({col}{sub_start_row}:{col}{sub_end_row}),2)"
        ws[f"{col}{low_mv_row}"] = f"=ROUND({col}{low_core_row}+{col}{add_total_row}-{col}{sub_total_row},2)"
        ws[f"{col}{high_mv_row}"] = f"=ROUND({col}{high_core_row}+{col}{add_total_row}-{col}{sub_total_row},2)"
        ws[f"{col}{low_px_row}"] = f"=ROUND(IF($B$7>0,{col}{low_mv_row}/$B$7,0),2)"
        ws[f"{col}{high_px_row}"] = f"=ROUND(IF($B$7>0,{col}{high_mv_row}/$B$7,0),2)"
        ws[f"{col}{low_px_row}"].number_format = "0.00"
        ws[f"{col}{high_px_row}"].number_format = "0.00"

    # Add two final rows to show converted lower/upper price limits in quote currency.
    lower_conv_row = r + 1
    upper_conv_row = r + 2
    ws.cell(lower_conv_row, 1, "Lower Price Limit (Quote Currency)")
    ws.cell(upper_conv_row, 1, "Upper Price Limit (Quote Currency)")
    ws.cell(lower_conv_row, len(years) + 2, f"{quote_currency}/股")
    ws.cell(upper_conv_row, len(years) + 2, f"{quote_currency}/股")
    for i in range(len(years)):
        col = get_column_letter(2 + i)
        ws[f"{col}{lower_conv_row}"] = f"=ROUND({col}{low_px_row}*$B$5,2)"
        ws[f"{col}{upper_conv_row}"] = f"=ROUND({col}{high_px_row}*$B$5,2)"
        ws[f"{col}{lower_conv_row}"].number_format = "0.00"
        ws[f"{col}{upper_conv_row}"].number_format = "0.00"

    wb.save(workbook_path)


def main() -> None:
    args = parse_args()
    code = _normalize_code(args.code)

    project_root = Path(__file__).resolve().parents[1]
    finance_root = project_root / "data" / "finance_report_web"

    income_df = pd.read_csv(finance_root / f"{code}_income_year.csv", index_col=0, encoding="utf-8-sig")
    balance_df = pd.read_csv(finance_root / f"{code}_balance_year.csv", index_col=0, encoding="utf-8-sig")
    cash_df = pd.read_csv(finance_root / f"{code}_cash_year.csv", index_col=0, encoding="utf-8-sig")
    info_df = pd.read_csv(finance_root / f"{code}_info.csv", index_col=0, encoding="utf-8-sig")

    company = args.company.strip()
    if not company and "股票名称" in info_df.index:
        try:
            company = str(info_df.loc["股票名称"].iloc[0]).strip()
        except Exception:
            company = ""
    if not company:
        company = code

    sheet_name = args.sheet.strip() or f"{company}_新估值"

    if args.years.strip().lower() == "auto":
        years = _infer_common_years(income_df, balance_df, cash_df, count=5)
    else:
        years = [int(x.strip()) for x in args.years.split(",") if x.strip()]

    if args.report_currency.strip():
        report_currency = _normalize_ccy(args.report_currency)
    else:
        report_currency = _detect_report_currency(income_df, balance_df, cash_df)
        if not _report_currency_marked(income_df, balance_df, cash_df):
            prev_ccy = _sheet_saved_report_currency(Path(args.workbook), sheet_name)
            if prev_ccy:
                report_currency = prev_ccy
    quote_currency = _normalize_ccy(args.quote_currency) if args.quote_currency.strip() else _quote_currency_from_code(
        code
    )

    fx_rate, fx_source = _resolve_fx_rate(
        workbook_path=Path(args.workbook),
        report_currency=report_currency,
        quote_currency=quote_currency,
        manual_fx=args.fx_rate,
    )

    rows, summary = build_sheet_rows(
        income_df_raw=income_df,
        balance_df_raw=balance_df,
        cash_df_raw=cash_df,
        info_df=info_df,
        code=code,
        years=years,
        fx_rate=fx_rate,
        yahoo_timeout=args.yahoo_timeout,
        shares_fallback_workbook=Path(args.workbook),
    )
    write_sheet(
        workbook_path=Path(args.workbook),
        sheet_name=sheet_name,
        company=company,
        years=years,
        rows=rows,
        summary=summary,
        overwrite=args.overwrite,
        report_currency=report_currency,
        quote_currency=quote_currency,
        fx_rate=fx_rate,
        fx_source=fx_source,
    )
    print(
        f"Wrote sheet '{sheet_name}' to {args.workbook} | years={years} "
        f"| report={report_currency} quote={quote_currency} fx={fx_rate} ({fx_source}) "
        f"| shares={summary.get('shares_source', '')}"
    )


if __name__ == "__main__":
    main()
