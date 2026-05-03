"""Export research reports to Word (.docx) format."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def markdown_report_to_docx(report_markdown: str, report_json: dict[str, Any]) -> BytesIO:
    """Convert a markdown research report into a styled .docx BytesIO buffer."""

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)

    for level in range(1, 4):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "Calibri"
        hs.font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)

    _add_title_page(doc, report_json)

    for line in report_markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            doc.add_heading(heading_match.group(2), level=level)
            continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet_match:
            _add_bullet(doc, bullet_match.group(1))
            continue

        doc.add_paragraph(stripped)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _add_title_page(doc: Document, report_json: dict[str, Any]) -> None:
    """Add a simple title page with company name, ticker, and rating."""

    company = report_json.get("company_name", "")
    ticker = report_json.get("ticker", "")
    rating = report_json.get("final_rating_label", report_json.get("final_rating", ""))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    run = title.add_run("Investment Research Report")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"{company} ({ticker})")
    sub_run.font.size = Pt(16)
    sub_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    if rating:
        rating_para = doc.add_paragraph()
        rating_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rating_para.paragraph_format.space_before = Pt(20)
        r = rating_para.add_run(f"Rating: {rating}")
        r.bold = True
        r.font.size = Pt(14)
        color_map = {"Buy": RGBColor(0x1B, 0x8A, 0x4A), "Watch": RGBColor(0xB8, 0x89, 0x16), "Avoid": RGBColor(0xC5, 0x4B, 0x4B)}
        r.font.color.rgb = color_map.get(rating, RGBColor(0x66, 0x66, 0x66))

    doc.add_page_break()


def _add_bullet(doc: Document, text: str) -> None:
    """Add a bullet paragraph, handling **bold** spans."""

    para = doc.add_paragraph(style="List Bullet")
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        else:
            para.add_run(part)
